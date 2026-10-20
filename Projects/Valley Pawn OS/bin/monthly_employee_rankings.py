#!/usr/bin/env python3
"""monthly_employee_rankings.py [YYYY-MM] [--render] [--no-pull] — native `monthly-employee-sales-rankings`
(1st of month 02:00 ET; Mac-first wave 3, 2026-10-19). The SKILL's steps, nothing re-derived:
  1 target = prior calendar month; pull `employee-activity-range` FIRST..LAST for all 5 stores through
    bin/bravo_pull.sh unless a full-month file (written after the month closed) is already on disk;
    one focused retry for the stores still missing
  2 parse: Employee rows + `Retail Sales Excluding Fees` (never another column); consolidate by name
  3b CURRENT EMPLOYEES ONLY (Joshua 2026-10-01): hr/ROSTER.json (Gusto active, <=4 days old) via
    comms_engine.load_active_keys / is_current; drop Preston Peters, service logins, former staff, $0 rows
  4 Rule 18: all 5 stores present + parsed + roster read, or nothing is posted (ledger row)
  5 #employee-performance: duplicate guard ("FINAL" + "Month YYYY" in the last 20 messages), main post
    (SKILL layout, same shape as the weekly board) + thread reply by store (FULL unfiltered store totals)
  6 workbook Employee Sales Rankings/Employee_Sales_Rankings_<Month>_<YYYY>.xlsx (never overwrites, _v2)
  7 working file monthly-analytics/<YYYY-MM> Employee Rankings.md (dropped names + reasons live here)
--render prints the post + thread and writes the workbook/working file to fleet/test_output/ only.
"""
import calendar
import datetime as dt
import os
import subprocess
import sys
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import comms_engine as ce   # noqa: E402  (read_csv_rows, money, usd, roster filter — the weekly board's own code)
import vp_slack             # noqa: E402

AGENT = "monthly-employee-sales-rankings"
ET = ZoneInfo("America/New_York")
OS_DIR = os.path.dirname(HERE)
BRAVO = os.environ.get("VP_BRAVO_DIR") or os.path.expanduser("~/Documents/Claude/Projects/Bravo Data Extraction")
LEDGER = os.path.join(OS_DIR, "fleet", "FAILURE_LEDGER.md")
CH = "C0ATTLPQHR8"   # #employee-performance
STORES = [("CUL", "Culpeper"), ("HAR", "Harrisonburg"), ("LEX", "Lexington"), ("ROA", "Roanoke"), ("WAY", "Waynesboro")]
THREAD_ORDER = ["CUL", "HAR", "ROA", "WAY", "LEX"]   # the live posts' by-store order
METRIC = "Retail Sales Excluding Fees"
PY = "/usr/bin/python3"


def ledger(sentence):
    try:
        open(LEDGER, "a").write("| %s (native) | %s | %s | NEEDS_HUMAN: no | OPEN |\n"
                                % (dt.datetime.now(ET).strftime("%Y-%m-%d %H:%M ET"), AGENT, sentence))
    except OSError:
        pass


def month_arg(args):
    if args:
        y, m = map(int, args[0].split("-"))
    else:
        first = dt.datetime.now(ET).date().replace(day=1) - dt.timedelta(days=1)
        y, m = first.year, first.month
    return dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])


def path(code, end):
    return os.path.join(BRAVO, "output", "%s_%s_employee-activity-range.csv" % (end, code))


def full_month_on_disk(code, start, end):
    """A usable file = present, non-trivial, written after the month closed, Reporting Dates = the exact month."""
    p = path(code, end)
    if not os.path.exists(p) or os.path.getsize(p) < 200:
        return False
    if dt.datetime.fromtimestamp(os.path.getmtime(p)).date() <= end:
        return False
    rows = ce.read_csv_rows(p)
    want = "%d/%d/%d - %d/%d/%d" % (start.month, start.day, start.year, end.month, end.day, end.year)
    return any("Reporting Dates" in (r[2] if len(r) > 2 else "") and want in ",".join(r) for r in rows[:6])


def pull(codes, start, end, tag):
    subprocess.run(["/bin/bash", os.path.join(HERE, "bravo_pull.sh"), "employee-activity-range",
                    "%s..%s" % (start, end), ",".join(codes), "monthly-emp-rankings-%s-%s-%d" % (start.strftime("%Y-%m"), tag, int(dt.datetime.now().timestamp()))],
                   capture_output=True, text=True, timeout=7500)


def parse(start, end):
    """-> (by_name {NAME: {total, stores, cols{}}}, store_totals {code: full unfiltered sum}, extra_cols, problems)"""
    by_name, store_totals, problems, extra_cols = {}, {}, [], []
    for code, name in STORES:
        p = path(code, end)
        if not os.path.exists(p):
            problems.append("%s: no full-month file" % name)
            continue
        rows = ce.read_csv_rows(p)
        hdr_i = next((i for i, r in enumerate(rows) if r and r[0].strip() == "Employee"), None)
        if hdr_i is None:
            problems.append("%s: no Employee header row" % name)
            continue
        hdr = [h.strip() for h in rows[hdr_i]]
        if METRIC not in hdr:
            problems.append("%s: '%s' column missing" % (name, METRIC))
            continue
        col = hdr.index(METRIC)
        extra_cols = [h for h in hdr[1:] if h and h != METRIC]
        n, tot = 0, 0.0
        for r in rows[hdr_i + 1:]:
            if not r or not r[0].strip():
                continue
            label = r[0].strip()
            if label.lower().startswith(("total store", "report printed")):
                continue
            emp = label.split(" - ", 1)[1].strip() if " - " in label else label
            try:
                v = ce.money(r[col]) if col < len(r) else 0.0
            except ValueError:
                continue
            tot += v   # EVERY login, unfiltered (SKILL step 3)
            key = emp.upper()
            d = by_name.setdefault(key, {"total": 0.0, "stores": [], "cols": {}})
            d["total"] += v
            if v != 0.0 and code not in d["stores"]:
                d["stores"].append(code)
            for h in extra_cols:
                j = hdr.index(h)
                try:
                    d["cols"][h] = d["cols"].get(h, 0.0) + (ce.money(r[j]) if j < len(r) and r[j].strip() else 0.0)
                except ValueError:
                    pass
            n += 1
        if n == 0:
            problems.append("%s: no employee rows" % name)
            continue
        store_totals[code] = tot
    return by_name, store_totals, extra_cols, problems


def classify(by_name, keys):
    ranked, dropped = [], []
    for k, v in by_name.items():
        if "PRESTON PETERS" in k:
            dropped.append((ce.title_name(k), "ownership")); continue
        if k.startswith(ce.SERVICE_LOGIN_PREFIXES):
            dropped.append((ce.title_name(k), "shared / service login")); continue
        if not ce.is_current(k, keys):
            dropped.append((ce.title_name(k), "not on the current team")); continue
        if v["total"] <= 0.0:
            dropped.append((ce.title_name(k), "$0.00")); continue
        ranked.append((k, v))
    ranked.sort(key=lambda kv: -kv[1]["total"])
    return ranked, dropped


def build(start, end, ranked, store_totals):
    month = start.strftime("%B %Y")
    medals = {1: ":first_place_medal:", 2: ":second_place_medal:", 3: ":third_place_medal:"}
    lines = ["*FINAL Employee Sales Rankings — Retail Sales Excluding Fees*",
             ":bar_chart: *%s* — full month, %d/1–%d/%d" % (month, start.month, end.month, end.day), ""]
    order = [c for c, _ in STORES]
    for i, (k, v) in enumerate(ranked, 1):
        stores = "+".join(sorted(v["stores"], key=order.index))
        name = ce.title_name(k)
        if i <= 3:
            lines.append("%s *%s* (%s) — %s" % (medals[i], name, stores, ce.usd(v["total"])))
        else:
            n = i
            suf = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
            lines.append("%d%s _%s_ (%s) — %s" % (n, suf, name, stores, ce.usd(v["total"])))
    company = sum(store_totals.values())
    lines += ["", "Company Total: %s" % ce.usd(company)]
    names = dict(STORES)
    thread = ["*:bar_chart: %s by store — Retail Sales Excluding Fees*" % month]
    thread += ["• %s: %s" % (names[c], ce.usd(store_totals[c])) for c in THREAD_ORDER]
    return "\n".join(lines), "\n".join(thread), company


def workbook(dest, start, ranked, dropped, by_name, store_totals, extra_cols):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    hdr_font, hdr_fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2D1A5E")
    fills = {1: "FFD700", 2: "C0C0C0", 3: "CD7F32"}
    order = [c for c, _ in STORES]

    def head(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill

    ws = wb.active
    ws.title = "Sales Rankings"
    head(ws, ["Rank", "Employee", "Store(s)", METRIC] + extra_cols)
    for i, (k, v) in enumerate(ranked, 1):
        ws.append([i, ce.title_name(k), "+".join(sorted(v["stores"], key=order.index)), round(v["total"], 2)]
                  + [round(v["cols"].get(h, 0.0), 2) for h in extra_cols])
        if i in fills:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=fills[i])
    ws2 = wb.create_sheet("All Logins (unfiltered)")
    head(ws2, ["Employee", "Store(s)", "Status", METRIC])
    status = {n: r for n, r in dropped}
    for k, v in sorted(by_name.items(), key=lambda kv: -kv[1]["total"]):
        nm = ce.title_name(k)
        st = {"ownership": "Ownership", "shared / service login": "Service", "not on the current team": "Former"}.get(status.get(nm), "Current")
        ws2.append([nm, "+".join(sorted(v["stores"], key=order.index)), st, round(v["total"], 2)])
    ws3 = wb.create_sheet("By Store")
    head(ws3, ["Store", METRIC + " (all logins)"])
    for c, n in STORES:
        ws3.append([n, round(store_totals.get(c, 0.0), 2)])
    ws3.append(["Company", round(sum(store_totals.values()), 2)])
    ws4 = wb.create_sheet("Summary")
    head(ws4, ["Item", "Value"])
    ws4.append(["Month", start.strftime("%B %Y")])
    ws4.append(["Company total (all logins)", round(sum(store_totals.values()), 2)])
    ws4.append(["Current employees ranked", len(ranked)])
    for i, (k, v) in enumerate(ranked[:3], 1):
        ws4.append(["#%d" % i, "%s — %s" % (ce.title_name(k), ce.usd(v["total"]))])
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
                if isinstance(c.value, float):
                    c.number_format = '"$"#,##0.00'
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    if os.path.exists(dest):
        dest = dest.replace(".xlsx", "_v2.xlsx")
    wb.save(dest)
    return dest


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    render, no_pull = "--render" in sys.argv, "--no-pull" in sys.argv
    start, end = month_arg(args)
    ym, month = start.strftime("%Y-%m"), start.strftime("%B %Y")
    notes = []
    missing = [c for c, _ in STORES if not full_month_on_disk(c, start, end)]
    if missing and not (render or no_pull):
        notes.append("pulled fresh for %s" % ", ".join(missing))
        pull(missing, start, end, "full")
        missing = [c for c, _ in STORES if not full_month_on_disk(c, start, end)]
        if missing:
            pull(missing, start, end, "retry-1")
            missing = [c for c, _ in STORES if not full_month_on_disk(c, start, end)]
    else:
        notes.append("reused full-month files already on disk" if not missing else "render/no-pull: missing %s" % ", ".join(missing))
    by_name, store_totals, extra_cols, problems = parse(start, end)
    problems = ["%s: file is not the full month (or written before the month closed)" % dict(STORES)[c] for c in missing] + problems
    roster_note = "hr/ROSTER.json"
    try:
        keys = ce.load_active_keys()
    except ce.Withhold as e:
        keys, problems = None, problems + list(e.args[0])
    out_dir = os.path.join(OS_DIR, "fleet", "test_output", "monthly_employee_rankings") if render else OS_DIR
    wf = os.path.join(out_dir, "monthly-analytics", "%s Employee Rankings.md" % ym)
    os.makedirs(os.path.dirname(wf), exist_ok=True)
    if problems:
        open(wf, "w").write("# %s Employee Sales Rankings — HELD (native)\n\nRun: %s\n\n%s\n\nNothing posted (Rule 18).\n"
                            % (month, dt.datetime.now(ET).strftime("%Y-%m-%d %H:%M ET"), "\n".join("- " + p for p in problems)))
        print("HOLD — " + "; ".join(problems))
        if not render:
            ledger("Employee sales rankings for %s held — %s." % (month, "; ".join(problems)[:300]))
        return 2
    ranked, dropped = classify(by_name, keys)
    if not ranked:
        print("HOLD — no current employees with sales")
        return 2
    post, thread, company = build(start, end, ranked, store_totals)
    xlsx = workbook(os.path.join(out_dir, "Employee Sales Rankings", "Employee_Sales_Rankings_%s.xlsx" % start.strftime("%B_%Y")),
                    start, ranked, dropped, by_name, store_totals, extra_cols)
    body = ["# %s Employee Sales Rankings — Working File (native)" % month, "",
            "Run: %s (%s)" % (dt.datetime.now(ET).strftime("%Y-%m-%d %H:%M ET"), "render" if render else "live"),
            "Metric: %s | Source files: %s_<STORE>_employee-activity-range.csv (%s)" % (METRIC, end, "; ".join(notes)),
            "Roster: %s" % roster_note, "",
            "## Dropped from the ranked list", ""] + ["- %s — %s" % d for d in sorted(dropped)] + ["",
            "## Ranked (current employees)", "", "| Rank | Employee | Store(s) | %s |" % METRIC, "|---|---|---|---|"]
    for i, (k, v) in enumerate(ranked, 1):
        body.append("| %d | %s | %s | %s |" % (i, ce.title_name(k), "+".join(v["stores"]), ce.usd(v["total"])))
    body += ["", "Company Total (all logins): %s" % ce.usd(company), "", "## By store (all logins)", ""]
    body += ["- %s: %s" % (dict(STORES)[c], ce.usd(store_totals[c])) for c in THREAD_ORDER]
    body += ["", "Workbook: %s" % xlsx, ""]
    if render:
        print("=== MAIN POST ===\n%s\n\n=== THREAD REPLY ===\n%s\n\n=== DROPPED ===\n%s\n\nworkbook: %s" % (post, thread, "\n".join("%s — %s" % d for d in sorted(dropped)), xlsx))
        open(wf, "w").write("\n".join(body))
        return 0
    os.environ["VP_TASK"] = AGENT
    try:   # SKILL step 5 duplicate guard: last 20 messages, "FINAL" + "Month YYYY"
        hist = vp_slack.call("conversations.history", params={"channel": CH, "limit": 20}).get("messages", [])
        dup = any("FINAL" in (m.get("text") or "") and month in (m.get("text") or "") for m in hist)
    except Exception:
        dup = False
    if dup:
        body.append("Not posted: a FINAL %s post already exists in the channel." % month)
        open(wf, "w").write("\n".join(body)); print("already posted for", month); return 0
    if vp_slack.dryrun_intercept("slack", CH, post + "\n\n--- thread ---\n" + thread):
        open(wf, "w").write("\n".join(body)); return 0
    r = vp_slack.call("chat.postMessage", {"channel": CH, "text": vp_slack.to_mrkdwn(post), "unfurl_links": False})
    if not r.get("ok"):
        ledger("Employee sales rankings for %s were built but did not post (%s)." % (month, r.get("error")))
        body.append("Not posted: %s" % r.get("error")); open(wf, "w").write("\n".join(body)); return 1
    vp_slack.receipt("slack", CH, True, len(post.encode()), post.splitlines()[0][:120])
    r2 = vp_slack.call("chat.postMessage", {"channel": CH, "text": vp_slack.to_mrkdwn(thread), "thread_ts": r["ts"], "unfurl_links": False})
    body.append("Posted: ts %s (thread reply %s)" % (r.get("ts"), "ok" if r2.get("ok") else r2.get("error")))
    open(wf, "w").write("\n".join(body))
    if not r2.get("ok"):
        ledger("Employee sales rankings for %s posted, but the by-store thread reply failed (%s)." % (month, r2.get("error")))
        return 1
    print("posted", month)
    return 0


if __name__ == "__main__":
    sys.exit(main())
