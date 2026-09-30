#!/usr/bin/env python3
"""
gift_credit_monthly.py — Valley Pawn monthly Gift Card & Store Credit report (deterministic).

Built 2026-09-29. Reads the Bravo pipeline CSVs for one closed month and produces:
  - <out>/gift_credit_<YYYY-MM>.json   machine-readable numbers + checks
  - <out>/gift_credit_<YYYY-MM>.xlsx   by-store summary + full detail tabs (customer names — internal only)
  - <out>/slack_post.txt               the Slack post, printed VERBATIM by the scheduled task

Inputs (Bravo Data Extraction/output/, produced by existing pipeline cells):
  <END>_<STORE>_credit-balance.csv   Bravo "Credit Balance", Ending Balance Date = <END> (single-date trigger)
  <END>_<STORE>_credit-journal.csv   Bravo "Credit Journal", <START>..<END>
  for STORE in CUL HAR LEX ROA WAY

What Bravo gives us (verified 2026-09-29 against the Aug 2026 pull):
  - Store Credits and Gift Cards are COMPANY-WIDE in Bravo (every store's report shows the same list);
    each row carries the ISSUING store, which is how we attribute by store.
  - Layaway Credits are PER STORE (each store's report shows only its own).
  - Credit Journal gives Beginning Balance / each transaction / Ending Balance per credit type, and the
    transaction's store prefix (VP4=CUL, VA5=HAR, VA1=LEX, ROA=ROA, VAP=WAY).

Exit codes: 0 = publish (checks pass) · 2 = HOLD (missing input or a check failed — publish nothing)
Rule 18: never post incomplete or inaccurate data. Every check must pass or nothing is posted.
"""
import csv, json, re, sys, os, datetime as dt
from collections import defaultdict

STORES = ["CUL", "HAR", "LEX", "ROA", "WAY"]
STORE_NAMES = {"CUL": "Culpeper", "HAR": "Harrisonburg", "LEX": "Lexington", "ROA": "Roanoke",
               "WAY": "Waynesboro", "STA": "Staunton (closed)"}
PREFIX = {"VP4": "CUL", "VA5": "HAR", "VA1": "LEX", "ROA": "ROA", "VAP": "WAY"}
TOL = 0.011


def money(s):
    s = (s or "").strip().replace("$", "").replace(",", "")
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def read_rows(path):
    with open(path, encoding="latin-1", newline="") as f:
        return [[c.strip() for c in r] for r in csv.reader(f)]


def last_money(row):
    for c in reversed(row):
        v = money(c)
        if v is not None:
            return v
    return None


def parse_date(s):
    try:
        return dt.datetime.strptime(s.split()[0], "%m/%d/%Y").date()
    except Exception:
        return None


# ---------------------------------------------------------------- Credit Balance
def parse_balance(path):
    rows = read_rows(path)
    out = {"asof": None, "store_credit": [], "gift_card": [], "layaway_credit": [], "totals": {}}
    sec = None
    for r in rows:
        if not any(r):
            continue
        if len(r) > 5 and r[2].startswith("Ending Balance Date"):
            out["asof"] = parse_date(r[5]) or next((parse_date(c) for c in r if parse_date(c)), None)
            continue
        first = r[0]
        if first == "Store Credits":
            sec = "store_credit"; continue
        if first == "Gift Cards":
            sec = "gift_card"; continue
        if first == "Layaway Credits":
            sec = "layaway_credit"; continue
        if first in ("Customer", "Gift Card Number", "Layaway Number"):
            continue
        if first.startswith("Total:"):
            if sec:
                out["totals"][sec] = last_money(r)
            continue
        if first.startswith("Report printed") or sec is None or first == "":
            continue
        if sec == "store_credit":
            # Customer,,Issued Date,Issuing Store,Remaining Amount
            out[sec].append({"customer": r[0], "issued": str(parse_date(r[2])), "store": r[3],
                             "original": None, "remaining": money(r[4])})
        else:
            # Number,Purchased By,Issued Date,Issuing Store,Original Amount,Remaining Amount
            out[sec].append({"number": r[0], "customer": r[1], "issued": str(parse_date(r[2])),
                             "store": r[3], "original": money(r[4]), "remaining": money(r[5])})
    return out


# ---------------------------------------------------------------- Credit Journal
def parse_journal(path):
    rows = read_rows(path)
    names = {"Store Credit": "store_credit", "Gift Card Credit": "gift_card", "Layaway Credit": "layaway_credit"}
    out = {k: {"begin": None, "end": None, "total": None, "txns": []} for k in names.values()}
    out["range"] = None
    sec = None
    for r in rows:
        if not any(r):
            continue
        r = r + [""] * (9 - len(r))
        if len(r) > 2 and r[2].startswith("Reporting Dates"):
            out["range"] = next((c for c in r[3:] if c), None)
            continue
        if r[0] in names:
            sec = names[r[0]]; continue
        if sec is None or r[0].startswith("Report printed"):
            continue
        if r[0].startswith("Beginning"):
            out[sec]["begin"] = last_money(r); continue
        if r[0].startswith("Ending"):
            out[sec]["end"] = last_money(r); continue
        if "Total:" in r:
            out[sec]["total"] = last_money(r); continue
        if r[0] in ("Customer", "Gift Card Number", "Layaway Number", "No data available"):
            continue
        # store column is index 4 in all three layouts
        store_raw = r[4] if len(r) > 4 else ""
        store = PREFIX.get(store_raw, store_raw)
        out[sec]["txns"].append({"id": r[0], "col1": r[1], "col2": r[2], "when": r[3], "store": store,
                                 "employee": r[5], "amount": money(r[8])})
    # Bravo omits the Beginning line (and leaves Total blank) when a credit type had no activity
    # in the period: the balance simply carried through, so begin = end and activity = 0.
    for k in names.values():
        s = out[k]
        if s["total"] is None and not s["txns"]:
            s["total"] = 0.0
        if s["begin"] is None and not s["txns"] and s["end"] is not None:
            s["begin"] = s["end"]
    return out


def main():
    if len(sys.argv) < 4:
        print("usage: gift_credit_monthly.py <YYYY-MM> <bravo_output_dir> <out_dir>", file=sys.stderr)
        sys.exit(2)
    ym, src, outdir = sys.argv[1], sys.argv[2], sys.argv[3]
    y, m = map(int, ym.split("-"))
    start = dt.date(y, m, 1)
    end = (dt.date(y + (m == 12), (m % 12) + 1, 1) - dt.timedelta(days=1))
    os.makedirs(outdir, exist_ok=True)
    problems = []

    bal, jr = {}, {}
    for s in STORES:
        bp = os.path.join(src, f"{end}_{s}_credit-balance.csv")
        jp = os.path.join(src, f"{end}_{s}_credit-journal.csv")
        for p in (bp, jp):
            if not os.path.exists(p) or os.path.getsize(p) < 300:
                problems.append(f"missing or empty input: {os.path.basename(p)}")
        if os.path.exists(bp):
            bal[s] = parse_balance(bp)
            if bal[s]["asof"] != end:
                problems.append(f"{s} credit-balance is as of {bal[s]['asof']}, expected {end}")
        if os.path.exists(jp):
            jr[s] = parse_journal(jp)
            want = f"{start.month}/{start.day}/{start.year} - {end.month}/{end.day}/{end.year}"
            if (jr[s]["range"] or "").replace(" ", "") != want.replace(" ", ""):
                problems.append(f"{s} credit-journal range is '{jr[s]['range']}', expected '{want}'")

    if problems:
        return hold(outdir, ym, problems)

    # --- company-wide sections must be identical in every store's file
    for sec in ("store_credit", "gift_card"):
        tots = {s: bal[s]["totals"].get(sec) for s in STORES}
        if len({round(v or -1, 2) for v in tots.values()}) != 1:
            problems.append(f"{sec} balance total differs between store files: {tots}")
        ends = {s: jr[s][sec]["end"] for s in STORES}
        if len({round(v or -1, 2) for v in ends.values()}) != 1:
            problems.append(f"{sec} journal ending balance differs between store files: {ends}")

    ref = bal["CUL"]
    rj = jr["CUL"]
    result = {"month": ym, "as_of": str(end), "generated": dt.datetime.now().isoformat(timespec="seconds"),
              "by_store": {}, "company": {}, "checks": []}

    def chk(name, a, b):
        ok = a is not None and b is not None and abs(a - b) <= TOL
        result["checks"].append({"check": name, "a": a, "b": b, "ok": ok})
        if not ok:
            problems.append(f"check failed: {name} ({a} vs {b})")

    # --- reconciliation checks
    for sec in ("store_credit", "gift_card"):
        rowsum = round(sum(x["remaining"] or 0 for x in ref[sec]), 2)
        chk(f"{sec}: detail rows = report total", rowsum, ref["totals"].get(sec))
        chk(f"{sec}: balance report = journal ending", ref["totals"].get(sec), rj[sec]["end"])
        j = rj[sec]
        chk(f"{sec}: journal begin + activity = ending", round((j["begin"] or 0) + (j["total"] or 0), 2), j["end"])
        chk(f"{sec}: journal txn sum = journal total",
            round(sum(t["amount"] or 0 for t in j["txns"]), 2), j["total"] if j["txns"] else 0.0)
    for s in STORES:
        b, j = bal[s], jr[s]["layaway_credit"]
        rowsum = round(sum(x["remaining"] or 0 for x in b["layaway_credit"]), 2)
        chk(f"{s} layaway credit: detail rows = report total", rowsum, b["totals"].get("layaway_credit"))
        chk(f"{s} layaway credit: balance report = journal ending", b["totals"].get("layaway_credit"), j["end"])
        chk(f"{s} layaway credit: journal begin + activity = ending",
            round((j["begin"] or 0) + (j["total"] or 0), 2), j["end"])

    # --- by-store figures
    all_keys = STORES + sorted({x["store"] for sec in ("store_credit", "gift_card") for x in ref[sec]} - set(STORES))
    two_years_ago = end.replace(year=end.year - 2)
    for s in all_keys:
        d = {"name": STORE_NAMES.get(s, s)}
        for sec in ("store_credit", "gift_card"):
            rows = [x for x in ref[sec] if x["store"] == s and (x["remaining"] or 0) > 0]
            d[f"{sec}_outstanding"] = round(sum(x["remaining"] for x in rows), 2)
            d[f"{sec}_count"] = len(rows)
            old = [x for x in rows if x["issued"] and x["issued"] != "None" and dt.date.fromisoformat(x["issued"]) < two_years_ago]
            d[f"{sec}_over_2yrs"] = round(sum(x["remaining"] for x in old), 2)
            txns = [t for t in rj[sec]["txns"] if t["store"] == s]
            d[f"{sec}_issued"] = round(sum(t["amount"] for t in txns if (t["amount"] or 0) > 0), 2)
            d[f"{sec}_redeemed"] = round(-sum(t["amount"] for t in txns if (t["amount"] or 0) < 0), 2)
        if s in bal:
            lrows = [x for x in bal[s]["layaway_credit"] if (x["remaining"] or 0) > 0]
            d["layaway_credit_outstanding"] = round(sum(x["remaining"] for x in lrows), 2)
            d["layaway_credit_count"] = len(lrows)
            lt = jr[s]["layaway_credit"]["txns"]
            d["layaway_credit_issued"] = round(sum(t["amount"] for t in lt if (t["amount"] or 0) > 0), 2)
            d["layaway_credit_redeemed"] = round(-sum(t["amount"] for t in lt if (t["amount"] or 0) < 0), 2)
            d["layaway_credit_begin"] = jr[s]["layaway_credit"]["begin"]
        else:
            for k in ("outstanding", "count", "issued", "redeemed"):
                d[f"layaway_credit_{k}"] = 0
        d["total_outstanding"] = round(d["store_credit_outstanding"] + d["gift_card_outstanding"]
                                       + d["layaway_credit_outstanding"], 2)
        result["by_store"][s] = d

    # store attribution must sum to the company totals
    for sec in ("store_credit", "gift_card"):
        chk(f"{sec}: by-store sum = company total",
            round(sum(v[f"{sec}_outstanding"] for v in result["by_store"].values()), 2),
            round(sum(x["remaining"] or 0 for x in ref[sec] if (x["remaining"] or 0) > 0), 2))
    unmapped = {t["store"] for sec in ("store_credit", "gift_card") for t in rj[sec]["txns"]} - set(STORES)
    if unmapped:
        problems.append(f"journal transactions with unknown store prefix: {sorted(unmapped)}")

    c = result["company"]
    for sec in ("store_credit", "gift_card"):
        c[f"{sec}_begin"] = rj[sec]["begin"]; c[f"{sec}_end"] = rj[sec]["end"]
        c[f"{sec}_issued"] = round(sum(v[f"{sec}_issued"] for v in result["by_store"].values()), 2)
        c[f"{sec}_redeemed"] = round(sum(v[f"{sec}_redeemed"] for v in result["by_store"].values()), 2)
    c["layaway_credit_begin"] = round(sum(jr[s]["layaway_credit"]["begin"] or 0 for s in STORES), 2)
    c["layaway_credit_end"] = round(sum(jr[s]["layaway_credit"]["end"] or 0 for s in STORES), 2)
    c["layaway_credit_issued"] = round(sum(result["by_store"][s]["layaway_credit_issued"] for s in STORES), 2)
    c["layaway_credit_redeemed"] = round(sum(result["by_store"][s]["layaway_credit_redeemed"] for s in STORES), 2)
    c["total_begin"] = round(c["store_credit_begin"] + c["gift_card_begin"] + c["layaway_credit_begin"], 2)
    c["total_end"] = round(c["store_credit_end"] + c["gift_card_end"] + c["layaway_credit_end"], 2)
    c["gift_cards_under_1"] = sum(1 for x in ref["gift_card"] if 0 < (x["remaining"] or 0) < 1)

    if problems:
        return hold(outdir, ym, problems, result)

    json.dump(result, open(os.path.join(outdir, f"gift_credit_{ym}.json"), "w"), indent=2, default=str)
    write_xlsx(os.path.join(outdir, f"gift_credit_{ym}.xlsx"), result, ref, bal, jr)
    post = render_post(result, start)
    open(os.path.join(outdir, "slack_post.txt"), "w").write(post)
    print(post)
    return 0


def hold(outdir, ym, problems, result=None):
    with open(os.path.join(outdir, "HOLD.txt"), "w") as f:
        f.write(f"HOLD {ym} — nothing published\n" + "\n".join(problems) + "\n")
        if result:
            f.write(json.dumps(result.get("checks", []), indent=1, default=str))
    print("HOLD: " + "; ".join(problems), file=sys.stderr)
    return 2


def f0(v):
    return f"${v:,.0f}" if abs(v) >= 1000 else f"${v:,.2f}"


def render_post(r, start):
    mon = start.strftime("%B %Y")
    c = r["company"]
    L = [f"*Gift Cards & Store Credit — {mon}*",
         f"What customers can still spend with us as of {dt.date.fromisoformat(r['as_of']).strftime('%b %-d')}. Credit is shown under the store that issued it.",
         ""]
    L.append("*Outstanding by store*")
    for s, d in sorted(r["by_store"].items(), key=lambda kv: -kv[1]["total_outstanding"]):
        if d["total_outstanding"] <= 0:
            continue
        parts = []
        if d["gift_card_outstanding"]:
            parts.append(f"gift cards {f0(d['gift_card_outstanding'])} ({d['gift_card_count']})")
        if d["store_credit_outstanding"]:
            parts.append(f"store credit {f0(d['store_credit_outstanding'])} ({d['store_credit_count']})")
        if d["layaway_credit_outstanding"]:
            parts.append(f"layaway credit {f0(d['layaway_credit_outstanding'])} ({d['layaway_credit_count']})")
        L.append(f"• {d['name']} — *{f0(d['total_outstanding'])}* · " + " · ".join(parts))
    L.append(f"• *Company — {f0(c['total_end'])}* (was {f0(c['total_begin'])} at the start of the month)")
    L.append("")
    L.append("*This month*")
    for s in STORES:
        d = r["by_store"][s]
        iss = d["gift_card_issued"] + d["store_credit_issued"] + d["layaway_credit_issued"]
        red = d["gift_card_redeemed"] + d["store_credit_redeemed"] + d["layaway_credit_redeemed"]
        if iss == 0 and red == 0:
            L.append(f"• {d['name']} — no activity")
            continue
        bits = []
        if d["gift_card_issued"] or d["gift_card_redeemed"]:
            bits.append(f"gift cards +{f0(d['gift_card_issued'])} / −{f0(d['gift_card_redeemed'])}")
        if d["store_credit_issued"] or d["store_credit_redeemed"]:
            bits.append(f"store credit +{f0(d['store_credit_issued'])} / −{f0(d['store_credit_redeemed'])}")
        if d["layaway_credit_issued"] or d["layaway_credit_redeemed"]:
            bits.append(f"layaway credit +{f0(d['layaway_credit_issued'])} / −{f0(d['layaway_credit_redeemed'])}")
        L.append(f"• {d['name']} — " + " · ".join(bits))
    L.append("_+ issued · − used. Gift cards include store credit a store put on a card; layaway credit is what customers paid on expired layaways._")
    old = sum(d["gift_card_over_2yrs"] + d["store_credit_over_2yrs"] for d in r["by_store"].values())
    if old > 0:
        L.append("")
        L.append(f"{f0(old)} of gift card and store credit balances were issued more than 2 years ago.")
    return "\n".join(L)


def write_xlsx(path, r, ref, bal, jr):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return
    wb = Workbook()
    ws = wb.active; ws.title = "By Store"
    hdr = ["Store", "Gift card $", "Gift cards #", "Store credit $", "Store credits #", "Layaway credit $",
           "Layaway credits #", "Total outstanding $", "GC issued", "GC used", "SC issued", "SC used",
           "LC issued", "LC used", "GC+SC issued >2 yrs ago $"]
    ws.append(hdr)
    for s, d in r["by_store"].items():
        ws.append([d["name"], d["gift_card_outstanding"], d["gift_card_count"], d["store_credit_outstanding"],
                   d["store_credit_count"], d["layaway_credit_outstanding"], d["layaway_credit_count"],
                   d["total_outstanding"], d["gift_card_issued"], d["gift_card_redeemed"], d["store_credit_issued"],
                   d["store_credit_redeemed"], d["layaway_credit_issued"], d["layaway_credit_redeemed"],
                   round(d["gift_card_over_2yrs"] + d["store_credit_over_2yrs"], 2)])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws2 = wb.create_sheet("Checks")
    ws2.append(["Check", "A", "B", "OK"])
    for x in r["checks"]:
        ws2.append([x["check"], x["a"], x["b"], "YES" if x["ok"] else "NO"])
    for sec, title in (("gift_card", "Gift Cards"), ("store_credit", "Store Credits")):
        w = wb.create_sheet(title)
        w.append(["Number", "Customer", "Issued", "Issuing store", "Original", "Remaining"])
        for x in ref[sec]:
            w.append([x.get("number", ""), x["customer"], x["issued"], x["store"], x["original"], x["remaining"]])
    w = wb.create_sheet("Layaway Credits")
    w.append(["Store", "Layaway #", "Customer", "Issued", "Credit amount", "Remaining"])
    for s in STORES:
        for x in bal[s]["layaway_credit"]:
            w.append([s, x["number"], x["customer"], x["issued"], x["original"], x["remaining"]])
    w = wb.create_sheet("Month Activity")
    w.append(["Type", "Store", "Card/Layaway/Customer", "Col2", "Col3", "When", "Employee", "Amount"])
    seen = set()
    for s in STORES:
        for sec in ("gift_card", "store_credit", "layaway_credit"):
            if sec != "layaway_credit" and s != "CUL":
                continue
            for t in jr[s][sec]["txns"]:
                w.append([sec, t["store"], t["id"], t["col1"], t["col2"], t["when"], t["employee"], t["amount"]])
    wb.save(path)


if __name__ == "__main__":
    sys.exit(main())
