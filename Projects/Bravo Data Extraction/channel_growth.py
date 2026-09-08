#!/usr/bin/env python3
"""
channel_growth.py -- revenue growth by CHANNEL and SUB-CHANNEL (additive, net-new)
==================================================================================
Created 2026-09-07. Answers: "is our NON-GOLD inventory contributing growth?"

THE DECOMPOSITION (all three sum exactly to Net Revenue -- verified to the penny
against the Bravo Company KPI report for June 2026):

    Net Revenue
      |
      +-- LOAN channel      = Pawn Service Charges          (interest + fees + misc)
      |
      +-- INVENTORY channel = Retail Sales Gross Profit Amt (NON-GOLD / everything cased)
                            + Scrap Sales Gross Profit Amt  (the melt channel)

WHY THIS SCRIPT EXISTS SEPARATELY FROM yield_by_asset_class.py
--------------------------------------------------------------
The End-of-Month export gives ONE combined `Sales Revenue (Profit)` figure -- it
cannot split retail from scrap. Only Bravo's **Company KPI** report carries the
split, on these rows:

    Retail Sales Gross Profit Amt      <- non-gold retail
    Scrap Sales / Scrap Sales Item Cost / Scrap Sales Gross Profit Amt
    Pawn Service Charges
    Net Revenue (Excluding Scrap/Repairs)

Confirmed exact for June 2026: PSC 77,160.79 + Retail GP 100,243.04 + Scrap GP
52,690.45 = 230,094.28 Net Revenue, and Scrap Item Cost 42,490.99 == the EOM
`Refined (Cost of Sales)` figure to the cent. So EOM and KPI agree; KPI just
carries one more level of detail.

INPUT: output/<END_DATE>_ALL_company-kpis.xlsx, one per period pulled.
       Pull via the `company-kpis` pipeline cell (stores:["ALL"], one fetch,
       all 5 stores as columns). NOTE the filename is keyed on END DATE only --
       so never pull two different ranges that share an end date, or the second
       silently overwrites the first (see eom_validate.py for the same trap).

USAGE:
    python3 channel_growth.py                      # every KPI file on disk
    python3 channel_growth.py 2025-12-31 2026-08-31
"""
import os, sys, glob, re, json
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'output')
STORES = ['CUL', 'HAR', 'LEX', 'ROA', 'WAY']

# Row labels we care about -> short keys. Matched on startswith, first hit wins,
# and the MTD duplicates further down the sheet are ignored because of that.
ROWS = {
    'Pawn Service Charges': 'psc',
    'Retail Sales Gross Profit Amt': 'retail_gp',
    'Retail Sales Total Amt': 'retail_sales',
    'Retail Sales Item Cost': 'retail_cost',
    'Scrap Sales Gross Profit Amt': 'scrap_gp',
    'Scrap Sales Item Cost': 'scrap_cost',
    'Scrap Sales': 'scrap_sales',          # must come AFTER the two above
    'Loan Balance': 'loan_bal',
    'Inventory Balance': 'inv_bal',
    'Layaway Balance': 'lay_bal',
    'New Loans Written Amt': 'new_loans',
    'New Buys Written Amt': 'new_buys',
    'Inventory Turns Annually': 'turns',
}
ORDER = ['Pawn Service Charges', 'Retail Sales Gross Profit Amt',
         'Retail Sales Total Amt', 'Retail Sales Item Cost',
         'Scrap Sales Gross Profit Amt', 'Scrap Sales Item Cost', 'Scrap Sales',
         'Loan Balance', 'Inventory Balance', 'Layaway Balance',
         'New Loans Written Amt', 'New Buys Written Amt', 'Inventory Turns Annually']


def money(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('$', '').replace(',', '').replace('%', '')
    if s in ('', '-'):
        return 0.0
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()')
    try:
        return -float(s) if neg else float(s)
    except ValueError:
        return None


def parse_kpi(path):
    """Return {'range':(start,end), 'total':{...}, 'stores':{'CUL':{...},...}}"""
    ws = openpyxl.load_workbook(path, data_only=True).active
    MC = ws.max_column

    # date range
    start = end = None
    for r in range(1, 20):
        for c in range(1, MC + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            s = str(v).strip()
            if s == 'Start Date':
                nx = [ws.cell(r, cc).value for cc in range(c + 1, MC + 1)]
                nx = [x for x in nx if x is not None]
                if nx:
                    start = str(nx[0])[:10]
            elif s == 'End Date':
                nx = [ws.cell(r, cc).value for cc in range(c + 1, MC + 1)]
                nx = [x for x in nx if x is not None]
                if nx:
                    end = str(nx[0])[:10]

    # header row: find the column of Grand Total and of each store
    colmap = {}
    hdr = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, MC + 1):
            v = ws.cell(r, c).value
            if v is not None and 'Grand Total' in str(v):
                hdr = r
                colmap['TOTAL'] = c
                for cc in range(1, MC + 1):
                    hv = ws.cell(r, cc).value
                    if hv is None:
                        continue
                    m = re.match(r'\s*(Cul|Har|Lex|Roa|Way)\b', str(hv), re.I)
                    if m:
                        colmap[m.group(1).upper()] = cc
                break
        if hdr:
            break
    if not hdr:
        return None

    data = {k: {} for k in list(colmap)}
    for r in range(hdr + 1, ws.max_row + 1):
        lab = None
        for c in range(1, colmap['TOTAL']):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                lab = str(v).strip()
                break
        if not lab:
            continue
        for pat in ORDER:
            key = ROWS[pat]
            if lab.startswith(pat):
                for who, cc in colmap.items():
                    if key not in data[who]:
                        val = money(ws.cell(r, cc).value)
                        if val is not None:
                            data[who][key] = val
                break
    return {'range': (start, end), 'file': os.path.basename(path), 'data': data}


def netrev(d):
    return (d.get('psc', 0) or 0) + (d.get('retail_gp', 0) or 0) + (d.get('scrap_gp', 0) or 0)


def main():
    files = sorted(glob.glob(os.path.join(OUT, '*_ALL_company-kpis.xlsx')))
    if len(sys.argv) > 1:
        files = [f for f in files if any(a in os.path.basename(f) for a in sys.argv[1:])]
    periods = []
    for f in files:
        p = parse_kpi(f)
        if p:
            periods.append(p)
    if not periods:
        print('no company-kpis files found')
        return

    for p in periods:
        t = p['data']['TOTAL']
        print(f"\n=== {p['range'][0]} .. {p['range'][1]}   ({p['file']})")
        print(f"  Loan revenue (PSC)          ${t.get('psc',0):>12,.0f}")
        print(f"  Retail GP (non-gold)        ${t.get('retail_gp',0):>12,.0f}"
              f"   on sales ${t.get('retail_sales',0):>12,.0f}")
        print(f"  Scrap GP (gold/melt)        ${t.get('scrap_gp',0):>12,.0f}"
              f"   on sales ${t.get('scrap_sales',0):>12,.0f}")
        print(f"  {'-'*58}")
        print(f"  NET REVENUE                 ${netrev(t):>12,.0f}")

    json.dump([{'range': p['range'], 'file': p['file'], 'data': p['data']} for p in periods],
              open(os.path.join(OUT, 'channel_growth.json'), 'w'), indent=1)
    print(f"\nwrote output/channel_growth.json ({len(periods)} periods)")


if __name__ == '__main__':
    main()
