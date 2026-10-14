#!/usr/bin/env python3
"""
feedback_ingest.py — turn Preston's (or Joshua's) #pawn-walks feedback into data the
valuation engine uses the next morning.

What it reads : #pawn-walks (C0B8WR95N31) — top-level human messages AND thread replies under
                any post, last N days. A feedback message is Preston re-posting the flag lines
                with his note after " - " (bold or not), e.g.
                  • L · $375 · 25% · GLOCK MODEL 41 GEN 4, SERIAL NUMBER BYAT4… - *Used value around 500, okay pay amount.*
                Free-text replies that quote no flag line are kept as notes (verdict UNCLEAR).
What it writes: feedback_ledger.jsonl   one row per annotated flag (idempotent on msg ts + line)
                manager_comps.json      model-key -> value Preston stated (used as a trusted comp)
                item_overrides.json     (store, desc, paid) Preston cleared -> never re-flag that item
                calibration.json        flag accuracy by category / reason / valuation source
                ingest_state.json       last-seen ts, last-run counts (what the walk footer reports)
How it runs   : run_daily_intake.py calls run() at the top of every walk (fail-soft, 25 s budget).
                Manual: python3 feedback_ingest.py [--days 14] [--dry-run] [--from-text file.txt]
Stdlib only. Slack token = vp_slack.token() (Keychain) -> run_daily_intake._get_slack_token().
"""
from __future__ import annotations
import csv, glob, json, os, re, sys, time, datetime, collections

HERE      = os.path.dirname(os.path.abspath(__file__))
PW_DIR    = os.path.dirname(HERE)                                   # Pawn Walks/
BRAVO_OUT = os.environ.get("VP_BRAVO_OUTPUT",
            "/Users/joshuadavis/Documents/Claude/Projects/Bravo Data Extraction/output/")
VP_BIN    = os.environ.get("VP_OS_BIN",
            "/Users/joshuadavis/Documents/Claude/Projects/Valley Pawn OS/bin")
CHANNEL   = "C0B8WR95N31"
LEDGER    = os.path.join(HERE, "feedback_ledger.jsonl")
COMPS     = os.path.join(HERE, "manager_comps.json")
OVERRIDES = os.path.join(HERE, "item_overrides.json")
CALIB     = os.path.join(HERE, "calibration.json")
STATE     = os.path.join(HERE, "ingest_state.json")
CFG       = os.path.join(HERE, "feedback_rules.json")

BOT_LINE = re.compile(r"""^\s*[•·\-\*]?\s*(?P<kind>RL|L|B)\s*[·|]\s*\$?\s*(?P<paid>[\d,]+(?:\.\d+)?)\s*[·|]\s*(?P<margin>-?\d+)\s*%\s*[·|]\s*(?P<rest>.+?)\s*$""")
STORE_HDR = re.compile(r"^\s*\*?(?P<store>CUL|HAR|LEX|ROA|WAY)\s*:?\*?\s*:?\s*$")
NOTE_SPLIT = re.compile(r"\s+[-–—]\s+\*?(?P<note>[^*]+?)\*?\s*$")

# ── config / io ───────────────────────────────────────────────────────────────
def _cfg():
    try:
        return json.load(open(CFG))
    except Exception:
        return {}

def _jload(p, default):
    try:
        return json.load(open(p))
    except Exception:
        return default

def _jsave(p, obj):
    tmp = p + ".tmp"
    with open(tmp, "w") as f:
        json.dump(obj, f, indent=1, sort_keys=True)
    os.replace(tmp, p)

# ── text normalisation (mirrors the engine, kept local so this file stands alone) ─
def norm_desc(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").upper().replace("…", "").strip())

def model_key(desc: str) -> str | None:
    """Stable per-model key: text before the first comma, 'MODEL' dropped, serials dropped.
    'GLOCK MODEL 41 GEN 4, SERIAL NUMBER …' -> 'GLOCK 41 GEN 4'. None when there is no
    model-ish token (a digit-bearing token or a 3+ letter brand + descriptor)."""
    d = norm_desc(desc)
    if re.search(r"\d\s*(DWT|GRAMS?|\bG\b|OZT?)\b|\b\d{1,2}K[- ]?[YWR]?/?G\b|SILV-?925|STERLING", d):
        return None                                   # weight-priced metal: melt is exact
    d = re.split(r",", d, 1)[0]
    d = re.sub(r"\bSERIAL\s*(NUMBER|NO\.?|#)?\s*[A-Z0-9-]*", " ", d)
    d = re.sub(r"\bMODEL\b", " ", d)
    d = re.sub(r"[^A-Z0-9/.\- ]", " ", d)
    d = re.sub(r"\s+", " ", d).strip(" -")
    toks = d.split()
    if len(toks) < 2:
        return None
    if not any(any(c.isdigit() for c in t) for t in toks) and len(toks[0]) < 3:
        return None
    return d

# ── verdict classification ────────────────────────────────────────────────────
_POS = re.compile(r"\b(ALL GOOD|OK(AY)?( PAY)?|FINE|GOOD( PAY)?|NOT AN? OVER ?PAY|FAIR|CORRECT|LEGIT|HALF[ -]?BACK|PURCHASED FROM US|BOUGHT FROM US|WE SOLD|OUR (OWN )?ITEM)\b")
_NEG = re.compile(r"\b(OVER ?PAY|OVER ?PAID|TOO MUCH|TOO HIGH|PAID TOO|\$?\d+ OVER|OVER BY)\b")
_NOT_NEG = re.compile(r"\b(NOT|NO|ISN'?T|NOT AN?) (AN? )?OVER ?PAY")
_REASONS = [
    ("HALF_BACK",     r"HALF[ -]?BACK|PURCHASED FROM US|BOUGHT FROM US|WE SOLD|OUR (OWN )?ITEM|SOLD (IT )?TO (HIM|HER|THEM)"),
    ("RELOAN",        r"RE-?LOAN|RE-?WRIT|RENEW"),
    ("QUANTITY",      r"QUANTITY|\bQTY\b|\bX ?\d\b|\b\d+ (COINS|PIECES|PCS|UNITS|OUNCES|OZ)\b"),
    ("ACCESSORY",     r"OPTIC|SCOPE|BATTER|CHARGER|\bKIT\b|\bCASE\b|ATTACH|INCLUDES|COMES WITH|\bSET\b"),
    ("INTERNAL_COMP", r"BRAVO|USED SALE|WE SOLD (ONE|THESE)|SOLD HISTORY|LAST SOLD"),
    ("MARKET_VALUE",  r"USED VALUE|WORTH|VALUE (IS|AROUND|ABOUT|~)|AROUND \$?\d|ABOUT \$?\d|~\s?\$?\d|RETAIL"),
    ("DEPRECIATED",   r"DEPRECIAT|OLD(ER)? MODEL|OUTDATED|AGED"),
    ("FUNCTION",      r"BROKEN|DOESN'?T WORK|NOT WORKING|DAMAGE|CRACK|DEFECT|MISSING"),
]
_VALUE = re.compile(r"(?:USED VALUE|VALUE|WORTH|SALE|SOLD|AROUND|ABOUT|~|IS|FOR|AT)[^$\d]{0,12}\$?\s*(\d{2,6}(?:,\d{3})*)(?![%\d])")
_OVER_BY = re.compile(r"\$?\s*(\d{1,5})\s*OVER\b|OVER BY\s*\$?\s*(\d{1,5})")
_QTY = re.compile(r"(?:QUANTITY|QTY)\s*(?:OF|:)?\s*(\d{1,3})\b|\bX\s?(\d{1,3})\b|\b(\d{1,3})\s*(?:COINS|PIECES|PCS|UNITS|OUNCES|OZ)\b")

def classify(note: str, paid: float) -> dict:
    n = norm_desc(note)
    neg = bool(_NEG.search(n)) and not _NOT_NEG.search(n)
    pos = bool(_POS.search(n)) or bool(_NOT_NEG.search(n))
    if neg:
        verdict = "OVERPAY"
    elif pos:
        verdict = "NOT_OVERPAY"
    else:
        verdict = "UNCLEAR"
    reasons = [name for name, pat in _REASONS if re.search(pat, n)]
    over_by = None
    m = _OVER_BY.search(n)
    if m:
        over_by = float(m.group(1) or m.group(2))
    value = None
    for m in _VALUE.finditer(n):
        v = float(m.group(1).replace(",", ""))
        if re.match(r"\s*OVER\b", n[m.end():]):      # "$10 over" is an amount, not a value
            continue
        if v != paid and v >= 5:
            value = v
            break
    qty = None
    m = _QTY.search(n)
    if m:
        qty = int(next(g for g in m.groups() if g))
    # a stated value decides an otherwise-unclear note: margin >= 30% on it = fine, else overpay
    if value is not None and verdict == "UNCLEAR":
        verdict = "NOT_OVERPAY" if (value - paid) / value >= 0.30 else "OVERPAY"
    if over_by is not None and verdict == "UNCLEAR":
        verdict = "OVERPAY"
    return {"verdict": verdict, "reasons": reasons, "stated_value": value,
            "over_by": over_by, "qty": qty}

# ── parse a message into annotated flag lines ─────────────────────────────────
def parse_message(text: str) -> list[dict]:
    out, store = [], None
    text = (text or "").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = STORE_HDR.match(line)
        if m:
            store = m.group("store")
            continue
        m = BOT_LINE.match(line)
        if not m:
            continue
        rest = m.group("rest")
        note = None
        nm = NOTE_SPLIT.search(rest)
        if nm:
            note = nm.group("note").strip()
            rest = rest[: nm.start()].strip()
        desc = rest.strip(" *_").replace("…", "").strip()
        if not note:
            continue                     # a re-posted line with no note carries no feedback
        out.append({"store": store, "kind": m.group("kind"),
                    "paid": float(m.group("paid").replace(",", "")),
                    "margin": int(m.group("margin")), "desc_prefix": desc, "note": note})
    return out

# ── match a line to the real intake row (store + paid + description prefix) ───
def _intake_rows(date: datetime.date) -> list[dict]:
    ds = date.isoformat()
    rows = []
    for f in glob.glob(os.path.join(BRAVO_OUT, f"{ds}_to_{ds}_*_intake-detail.csv")) + \
             glob.glob(os.path.join(BRAVO_OUT, f"{ds}_*_intake-detail.csv")):
        m = re.search(r"_([A-Z]{3})_intake-detail", os.path.basename(f))
        st = m.group(1) if m else "UNK"
        try:
            with open(f, newline="", encoding="utf-8-sig") as fh:
                for r in csv.DictReader(fh):
                    amt = (r.get("Loan Amount") or r.get("Amount") or "").replace("$", "").replace(",", "").strip()
                    try:
                        amt = float(amt)
                    except ValueError:
                        continue
                    rows.append({"store": st, "date": ds, "paid": amt,
                                 "desc": norm_desc(r.get("Full Description") or r.get("FullDescription")),
                                 "category": (r.get("Category") or "").strip(),
                                 "kind": (r.get("Ticket Kind") or "").strip().upper()})
        except Exception:
            pass
    return rows

def match_line(line: dict, msg_date: datetime.date, lookback: int = 8) -> dict | None:
    pref = norm_desc(line["desc_prefix"])[:30]
    for back in range(1, lookback + 1):
        d = msg_date - datetime.timedelta(days=back)
        for r in _intake_rows(d):
            if abs(r["paid"] - line["paid"]) < 0.51 and r["desc"].startswith(pref) \
               and (line["store"] in (None, r["store"])):
                return r
    return None

# ── slack ─────────────────────────────────────────────────────────────────────
def _slack_call(method, params):
    try:
        sys.path.insert(0, VP_BIN)
        import vp_slack                                   # Keychain token, stdlib
        return vp_slack.call(method, params=params)
    except SystemExit:
        pass
    except Exception:
        pass
    # fallback: the walk's own token finder
    try:
        sys.path.insert(0, PW_DIR)
        import run_daily_intake as rdi                    # noqa
        tok = rdi._get_slack_token()
    except Exception:
        tok = None
    if not tok:
        raise RuntimeError("no slack token")
    import urllib.request, urllib.parse
    url = "https://slack.com/api/" + method + "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"Authorization": "Bearer " + tok})
    return json.load(urllib.request.urlopen(req, timeout=25))

def fetch_feedback_messages(days: int, since_ts: float = 0.0) -> list[dict]:
    """Human messages (top-level + thread replies) in #pawn-walks newer than since_ts."""
    # No "oldest" param on purpose: bin/vp_slack.py saw oldest-bounded history calls return an
    # empty page intermittently (2026-10-05). Read newest-first pages and stop once past the window.
    oldest = max(since_ts, time.time() - days * 86400)
    msgs, cursor = [], None
    for _ in range(5):
        p = {"channel": CHANNEL, "limit": 200}
        if cursor:
            p["cursor"] = cursor
        r = _slack_call("conversations.history", p)
        if not r.get("ok"):
            raise RuntimeError("conversations.history: " + str(r.get("error")))
        page = r.get("messages", [])
        msgs += [m for m in page if float(m.get("ts", 0)) >= oldest]
        cursor = (r.get("response_metadata") or {}).get("next_cursor")
        if not cursor or (page and float(page[-1].get("ts", 0)) < oldest):
            break
    out = []
    for m in msgs:
        if m.get("user") and not m.get("bot_id") and m.get("subtype") in (None, "thread_broadcast", "file_share"):
            out.append(m)
        if (m.get("reply_count") or 0) > 0:
            rr = _slack_call("conversations.replies", {"channel": CHANNEL, "ts": m["ts"], "limit": 200})
            for x in rr.get("messages", [])[1:]:
                if x.get("user") and not x.get("bot_id") and float(x.get("ts", 0)) > since_ts:
                    x["_parent_ts"] = m["ts"]
                    out.append(x)
    return [m for m in out if float(m.get("ts", 0)) > since_ts]

# ── derive durable artefacts from the ledger ──────────────────────────────────
def rebuild_artifacts(ledger: list[dict]):
    cfg = _cfg()
    comps = _jload(COMPS, {})
    overrides = _jload(OVERRIDES, {})
    today = datetime.date.today()
    for e in ledger:
        c = e.get("class") or {}
        key = e.get("model_key")
        # 1) manager comp: a stated value on a model we can key -> trusted comp
        if c.get("stated_value") and key:
            cur = comps.get(key)
            if not cur or cur.get("date", "") <= e["fb_date"]:
                comps[key] = {"value": c["stated_value"], "date": e["fb_date"], "by": e["by"],
                              "paid": e["paid"], "store": e.get("store"), "desc": e.get("desc"),
                              "note": e["note"]}
        # 2) cleared item: don't flag this exact item again (reloans of the same ticket)
        if c.get("verdict") == "NOT_OVERPAY" and e.get("desc"):
            okey = f"{e.get('store')}|{e['desc']}|{e['paid']:.0f}"
            overrides[okey] = {"date": e["fb_date"], "by": e["by"], "reasons": c.get("reasons", []),
                               "note": e["note"], "generic": (key is None) or ("SERIAL" not in (e.get("desc") or ""))}
    _jsave(COMPS, comps)
    _jsave(OVERRIDES, overrides)
    # 3) calibration
    win = (cfg.get("category_tuning") or {}).get("window_days", 90)
    cut = (today - datetime.timedelta(days=win)).isoformat()
    by_cat = collections.defaultdict(lambda: {"confirmed": 0, "false": 0, "unclear": 0})
    by_reason = collections.Counter()
    by_src = collections.defaultdict(lambda: {"confirmed": 0, "false": 0, "unclear": 0})
    tot = {"confirmed": 0, "false": 0, "unclear": 0}
    for e in ledger:
        if e["fb_date"] < cut:
            continue
        v = (e.get("class") or {}).get("verdict")
        b = "confirmed" if v == "OVERPAY" else "false" if v == "NOT_OVERPAY" else "unclear"
        by_cat[e.get("category") or "?"][b] += 1
        by_src[e.get("val_source") or "?"][b] += 1
        tot[b] += 1
        for r in (e.get("class") or {}).get("reasons", []):
            by_reason[f"{r}|{b}"] += 1
    tun = cfg.get("category_tuning") or {}
    tuned = {}
    for cat, d in by_cat.items():
        n = d["confirmed"] + d["false"]
        if d["false"] >= tun.get("min_false", 3) and n and d["false"] / n >= tun.get("false_rate", 0.6):
            tuned[cat] = {"flag_margin": tun.get("severe_flag_margin", 0.15),
                          "why": f"{d['false']} of {n} flags called wrong in last {win}d"}
    calib = {"as_of": today.isoformat(), "window_days": win, "totals": tot,
             "precision": (tot["confirmed"] / (tot["confirmed"] + tot["false"])) if (tot["confirmed"] + tot["false"]) else None,
             "by_category": dict(by_cat), "by_source": dict(by_src),
             "by_reason": dict(by_reason), "category_tuning": tuned,
             "ledger_rows": len(ledger)}
    _jsave(CALIB, calib)
    return comps, overrides, calib

# ── main ──────────────────────────────────────────────────────────────────────
def _load_ledger():
    rows = []
    try:
        for ln in open(LEDGER):
            ln = ln.strip()
            if ln:
                rows.append(json.loads(ln))
    except FileNotFoundError:
        pass
    return rows

def ingest_messages(msgs: list[dict], dry_run=False) -> dict:
    cfg = _cfg()
    names = cfg.get("reviewers", {})
    ledger = _load_ledger()
    seen = {(e["msg_ts"], e["line"]) for e in ledger}
    new, notes, unmatched = [], [], []
    for m in msgs:
        ts = m["ts"]
        msg_dt = datetime.datetime.fromtimestamp(float(ts))
        by = names.get(m.get("user"), m.get("user", "?"))
        lines = parse_message(m.get("text") or "")
        if not lines:
            txt = (m.get("text") or "").strip()
            if len(txt) > 12 and "has joined" not in txt and (ts, -1) not in seen:
                notes.append({"msg_ts": ts, "line": -1, "by": by, "fb_date": msg_dt.date().isoformat(),
                              "note": txt[:600], "class": {"verdict": "UNCLEAR", "reasons": []},
                              "matched": False})
            continue
        for i, ln in enumerate(lines):
            if (ts, i) in seen:
                continue
            row = match_line(ln, msg_dt.date())
            c = classify(ln["note"], ln["paid"])
            e = {"msg_ts": ts, "line": i, "by": by, "fb_date": msg_dt.date().isoformat(),
                 "store": (row or {}).get("store") or ln["store"], "kind": ln["kind"],
                 "paid": ln["paid"], "walk_margin": ln["margin"] / 100.0,
                 "desc": (row or {}).get("desc") or norm_desc(ln["desc_prefix"]),
                 "category": (row or {}).get("category"), "walk_date": (row or {}).get("date"),
                 "matched": row is not None, "note": ln["note"], "class": c,
                 "model_key": model_key((row or {}).get("desc") or ln["desc_prefix"]),
                 "val_source": None}
            # valuation source for the calibration table, from the walk's own spreadsheet
            e["val_source"] = _lookup_source(e)
            new.append(e)
            if row is None:
                unmatched.append(e)
    all_new = new + notes
    if not dry_run and all_new:
        with open(LEDGER, "a") as f:
            for e in all_new:
                f.write(json.dumps(e, sort_keys=True) + "\n")
        ledger += all_new
        rebuild_artifacts(ledger)
    return {"new_lines": len(new), "new_notes": len(notes), "unmatched": len(unmatched),
            "by": sorted({e["by"] for e in all_new}),
            "walk_dates": sorted({e["walk_date"] for e in new if e.get("walk_date")}),
            "verdicts": dict(collections.Counter(e["class"]["verdict"] for e in new)),
            "comps_added": sum(1 for e in new if e["class"].get("stated_value") and e.get("model_key")),
            "overrides_added": sum(1 for e in new if e["class"]["verdict"] == "NOT_OVERPAY"),
            "entries": all_new}

def _lookup_source(e):
    """Valuation source the walk used for this item (xlsx Items tab), for the calibration table."""
    if not e.get("walk_date"):
        return None
    p = os.path.join(PW_DIR, "daily", f"{e['walk_date']}_intake_margin.xlsx")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        hdr = None
        for row in ws.iter_rows(values_only=True):
            if hdr is None:
                if row and any(str(c).strip().lower() in ("desc", "description", "item") for c in row if c) \
                   and any(str(c).strip().lower() in ("cost paid", "cost", "paid") for c in row if c):
                    hdr = [str(c).strip().lower() if c else "" for c in row]
                continue
            d = dict(zip(hdr, row))
            desc = norm_desc(str(d.get("desc") or d.get("description") or d.get("item") or ""))
            cost = d.get("cost paid") or d.get("cost") or d.get("paid")
            try:
                cost = float(cost)
            except Exception:
                continue
            if desc[:30] == e["desc"][:30] and abs(cost - e["paid"]) < 0.51:
                return str(d.get("value source") or d.get("source") or d.get("src") or "")
    except Exception:
        return None
    return None

def run(days: int = 14, dry_run: bool = False, quiet: bool = True) -> dict:
    state = _jload(STATE, {})
    since = float(state.get("last_ts", 0) or 0)
    t0 = time.time()
    try:
        msgs = fetch_feedback_messages(days, since)
    except Exception as ex:
        res = {"ok": False, "error": str(ex)[:200], "new_lines": 0, "new_notes": 0}
        if not quiet:
            print("feedback_ingest:", res)
        return res
    res = ingest_messages(msgs, dry_run=dry_run)
    res["ok"] = True
    res["elapsed"] = round(time.time() - t0, 1)
    if not dry_run:
        if msgs:
            state["last_ts"] = max(float(m["ts"]) for m in msgs)
        state["last_run"] = datetime.datetime.now().isoformat(timespec="seconds")
        state["last_result"] = {k: v for k, v in res.items() if k != "entries"}
        _jsave(STATE, state)
    if not quiet:
        print(json.dumps({k: v for k, v in res.items() if k != "entries"}, indent=1))
        for e in res.get("entries", []):
            print(f"  [{e['class']['verdict']:<12}] {e.get('store')} ${e['paid']:.0f} {e.get('desc','')[:45]:<45} "
                  f"reasons={','.join(e['class'].get('reasons', [])) or '-'} value={e['class'].get('stated_value')} matched={e['matched']}")
    return res

if __name__ == "__main__":
    a = sys.argv[1:]
    days = int(a[a.index("--days") + 1]) if "--days" in a else 14
    if "--from-text" in a:                      # offline test: one message in a text file
        txt = open(a[a.index("--from-text") + 1], encoding="utf-8").read()
        ts = a[a.index("--ts") + 1] if "--ts" in a else f"{time.time():.6f}"
        user = a[a.index("--user") + 1] if "--user" in a else "U03BWMEM9GR"
        res = ingest_messages([{"ts": ts, "user": user, "text": txt}], dry_run="--dry-run" in a)
        print(json.dumps({k: v for k, v in res.items() if k != "entries"}, indent=1))
        for e in res["entries"]:
            print(f"  [{e['class']['verdict']:<12}] {e.get('store')} ${e['paid']:.0f} {e.get('desc','')[:45]:<45} "
                  f"reasons={','.join(e['class'].get('reasons', [])) or '-'} value={e['class'].get('stated_value')} "
                  f"qty={e['class'].get('qty')} key={e.get('model_key')} matched={e['matched']} src={e.get('val_source')}")
    elif "--rebuild" in a:
        print(json.dumps(rebuild_artifacts(_load_ledger())[2], indent=1))
    else:
        run(days=days, dry_run="--dry-run" in a, quiet=False)
