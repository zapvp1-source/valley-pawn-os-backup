#!/usr/bin/env python3
"""Training Time Is Paid Time — house-format policy (Oct 2026). Handbook-class (pay practice, FLSA 29 CFR 785.27)."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x2D, 0x1A, 0x5E)
BLUE = RGBColor(0x00, 0x99, 0xDD)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

OUT = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(OUT, "Training_At_The_Store_Policy_2026-10.docx")

TITLE = "Training Happens at the Store"
STARTS = "Starts October 5, 2026   ·   All stores"
BOX = ["Do all required training at the store,", "on the clock, at your normal rate."]
LINES = [
    "Train during slow time on your shift, during business hours.",
    "No training at home or off the clock.",
    "Running behind your due date? Tell your manager that day.",
]
RATE_NOTE = None
LEGAL = ("I received, understand, and will follow this policy. The company may change or end it at any time "
         "without notice. This is not an employment contract and does not change at-will employment; either "
         "the company or I may end my employment at any time, with or without cause or notice.")

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(0.5)
s.left_margin = s.right_margin = Inches(0.75)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(12)


def para(text, size, bold=False, color=None, after=0, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return p


para("VALLEY PAWN", 9, bold=True, color=BLUE, after=0)
para(TITLE, 22, bold=True, color=NAVY, after=0)
para(STARTS, 9.5, color=GREY, after=10)

box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.rows[0].cells[0]
cell.width = Inches(7.0)
sh = OxmlElement("w:shd")
sh.set(qn("w:val"), "clear")
sh.set(qn("w:fill"), "2D1A5E")
cell._tc.get_or_add_tcPr().append(sh)
cell.text = ""
for i, line in enumerate(BOX):
    p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10 if i == 0 else 0)
    p.paragraph_format.space_after = Pt(10 if i == len(BOX) - 1 else 2)
    r = p.add_run(line)
    r.font.size = Pt(17)
    r.bold = True
    r.font.color.rgb = WHITE

para("", 6, after=4)
for line in LINES:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    r = p.add_run("→  ")
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = BLUE
    r2 = p.add_run(line)
    r2.font.size = Pt(12)

if RATE_NOTE:
    para(RATE_NOTE, 9.5, color=GREY, before=2, after=0)
para(LEGAL, 8, color=GREY, before=10, after=6)

sig = doc.add_table(rows=2, cols=4)
sig.alignment = WD_TABLE_ALIGNMENT.LEFT
sw = [Inches(2.5), Inches(1.7), Inches(1.4), Inches(1.4)]
labels = ["Signature (via Gusto)", "Printed Name", "Store", "Date"]
for j, lab in enumerate(labels):
    c = sig.rows[0].cells[j]
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("_" * (22 if j == 0 else 16))
    r.font.size = Pt(10)
    c2 = sig.rows[1].cells[j]
    c2.text = ""
    r2 = c2.paragraphs[0].add_run(lab)
    r2.font.size = Pt(8)
    r2.font.color.rgb = GREY
for row in sig.rows:
    for j, w in enumerate(sw):
        row.cells[j].width = w

doc.save(DOCX)
print("wrote", DOCX)
