import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY=RGBColor(0x2D,0x1A,0x5E); BLUE=RGBColor(0x00,0x99,0xDD); GREY=RGBColor(0x59,0x59,0x59); WHITE=RGBColor(255,255,255)
LEGAL=("I received and understand this policy and will follow it. Valley Pawn may change or end it at any time, "
       "with or without notice. This is not a contract of employment and does not change at-will employment — "
       "either I or the company may end employment at any time, with or without cause or notice.")
START="Starts Monday, October 5, 2026   ·   All stores"

def base(doc):
    s=doc.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(0.75); s.top_margin=s.bottom_margin=Inches(0.5)
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    st.paragraph_format.space_after=Pt(0)

def para(doc_or_cell, text='', size=12, bold=False, color=None, align=None, before=0, after=0):
    p=doc_or_cell.add_paragraph()
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    if align: p.alignment=align
    if text:
        r=p.add_run(text); r.font.size=Pt(size); r.bold=bold
        if color is not None: r.font.color.rgb=color
    return p

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); tcPr.append(sh)

def borders(table, color='FFFFFF', sz=0):
    tbl=table._tbl; tblPr=tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+e); el.set(qn('w:val'),'single' if sz else 'nil'); el.set(qn('w:sz'),str(sz)); el.set(qn('w:color'),color); b.append(el)
    tblPr.append(b)

def cellmargins(cell, t=80,b=80,l=120,r=120):
    tcPr=cell._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for k,v in (('top',t),('bottom',b),('start',l),('end',r),('left',l),('right',r)):
        el=OxmlElement('w:'+k); el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa'); m.append(el)
    tcPr.append(m)

def header(doc, title):
    para(doc,'VALLEY PAWN',9,True,BLUE)
    para(doc,title,22,True,NAVY)
    para(doc,START,9.5,False,GREY,after=8)

def arrow(container, text, size=12, after=3):
    p=container.add_paragraph(); p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.left_indent=Inches(0.22); p.paragraph_format.first_line_indent=Inches(-0.22)
    r=p.add_run('→  '); r.bold=True; r.font.color.rgb=NAVY; r.font.size=Pt(size)
    r=p.add_run(text); r.font.size=Pt(size)
    return p

def signature(doc):
    t=doc.add_table(rows=2, cols=4); borders(t)
    for i,u in enumerate(['_'*20,'_'*16,'_'*16,'_'*16]):
        c=t.rows[0].cells[i]; c.paragraphs[0].add_run(u).font.size=Pt(11)
    for i,l in enumerate(['Signature (via Gusto)','Printed Name','Store','Date']):
        r=t.rows[1].cells[i].paragraphs[0].add_run(l); r.font.size=Pt(8.5); r.font.color.rgb=GREY

# ---------------- 1. signable policy ----------------
d=docx.Document(); base(d)
header(d,'Account Details by Phone & Text')
box=d.add_table(rows=1, cols=1); box.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(box)
c=box.rows[0].cells[0]; shade(c,'2D1A5E'); cellmargins(c,160,160,200,200)
p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Only the account holder hears account details —'); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=WHITE
p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('and only after THEY tell YOU their full name and date of birth.'); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=WHITE
para(d,'',after=6)
arrow(d,'Anyone else, spouse included, comes in with ID and the ticket, or the customer calls.')
arrow(d,'Never look up or confirm a customer for a caller — call police back on their department’s main line.')
arrow(d,'Said too much, or not sure you should? Tell Preston the same day.')
para(d,LEGAL,8,False,GREY,before=10,after=10)
signature(d)
d.save('/tmp/cw/scripts/Account_Details_Phone_Text_Policy_2026-10.docx')

# ---------------- 2. counter card ----------------
d=docx.Document(); base(d)
s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(0.4); s.left_margin=s.right_margin=Inches(0.5)
para(d,'VALLEY PAWN',9,True,BLUE)
para(d,'Phone & Text Scripts',22,True,NAVY)
para(d,START+'   ·   Keep at the counter',9.5,False,GREY,after=6)

CARDS=[
 ("1  ANSWERING",
  "“Valley Pawn [city], this is [name]. We loan, buy and sell — how can I help?”",
  ["Ask their name and use it.",
   "Ask before you put anyone on hold. Over a minute? Take their number and call back within 15 minutes."]),
 ("2  THEIR ACCOUNT",
  "“Before I pull that up, can I get your full name and date of birth?”",
  ["They say it. You never read it to them or ask “is it ___?”",
   "Only the person on the account. Spouse or anyone else: come in with ID and the ticket.",
   "Can’t verify: “Sorry, I can only go over an account with the account holder.”",
   "Phone number change: in store, with ID. Never by phone or text."]),
 ("3  ASKING ABOUT SOMEONE ELSE",
  "“I can’t say whether anyone has done business with us. If something was stolen, file a police report — we send police everything we take in, every day.”",
  ["Never search a name or ticket for a caller.",
   "Says they’re police? Get name, agency, badge and case number, then call back on the department’s main line.",
   "Tell Preston the same day."]),
 ("4  GUNS",
  "“Handguns are 21 and up, long guns 18 and up. No exceptions.”",
  ["A gun on an account: verify first (card 2).",
   "Background check cleared: call the number on file. Voicemail: “Please call Valley Pawn.” Nothing more.",
   "Pickup: only the buyer or the account holder, with ID. Background check at pickup.",
   "Gun we sold won’t work: “Bring it in unloaded and cased and the manager will take care of you.” No repair tips. Tell Preston."]),
 ("5  SELLING OR PAWNING",
  "“Have you done business with us before? Looking to sell it or borrow on it? How much are you hoping to get?”",
  ["Ask them to text photos to this number before they drive in.",
   "Offer too low to sell? Offer a loan. Short of what they need? Ask what else they have.",
   "We don’t take cell phones. Not sure about anything else? “Let me check and call you back today.”"]),
 ("6  LOAN PAYMENTS",
  "“Your loan was due [date from Bravo]. You can pay in the app or come in.”",
  ["Dates and amounts come from Bravo, never memory. Don’t quote grace periods or fees.",
   "“How long do I have?” — “Once it’s past due it can be pulled for sale, so let’s get it paid.”",
   "No card payments over the phone."]),
 ("7  TEXTS",
  "Answer every customer text the same day. After hours: first thing next morning.",
  ["Prices and “come see it” are fine by text. Accounts and guns: call them.",
   "Don’t close a conversation until the customer has an answer.",
   "“Stop” or “don’t contact me”: note it in Bravo and tell Preston."]),
 ("8  EVERY CALL",
  "Every “I’ll call you back” gets a call back the same day.",
  ["Voicemail: say it, then hang up before you talk to anyone else.",
   "Tax is the same, cash or card.",
   "Don’t have it? Check our other four stores before sending anyone elsewhere.",
   "Upset customer: stay calm, fix what you can, tell Preston the same day."]),
]
t=d.add_table(rows=4, cols=2); borders(t,'FFFFFF',0)
t.autofit=False
for i,(title,line,bul) in enumerate(CARDS):
    c=t.rows[i//2].cells[i%2]; c.width=Inches(3.75); cellmargins(c,70,70,110,110)
    shade(c,'F2F0F8')
    p=c.paragraphs[0]; r=p.add_run(title); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=BLUE
    p=c.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(line); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=NAVY
    for b in bul: arrow(c,b,9.5,1)
# gutter spacing between cells
tblPr=t._tbl.tblPr; cs=OxmlElement('w:tblCellSpacing'); cs.set(qn('w:w'),'60'); cs.set(qn('w:type'),'dxa'); tblPr.append(cs)
para(d,'Questions about any of these? Ask Preston. The account-details rule is signed separately in Gusto.',8,False,GREY,before=4)
d.save('/tmp/cw/scripts/Phone_Text_Scripts_Counter_Card_2026-10.docx')
print('built')
