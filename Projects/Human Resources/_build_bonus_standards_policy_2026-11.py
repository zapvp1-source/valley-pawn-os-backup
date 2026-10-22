#!/usr/bin/env python3
"""Bonus Standards & Management Rights — house-format policy (Oct 2026, rides PU-2026-11). Handbook-class (pay practice; Va. Code 40.1-29, 29 CFR 778.211). Text mirrors Bonus Program/bonus_rules.json management_rights."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x2D, 0x1A, 0x5E)
BLUE = RGBColor(0x00, 0x99, 0xDD)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

OUT = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(OUT, "Bonus_Standards_Management_Rights_Policy_2026-11.docx")

TITLE = "Bonuses — Standards First"
STARTS = "Starts with the December 2026 bonus month   ·   All stores"
BOX = ["No bonus unless your store meets", "every company standard that month."]
LINES = [
    "Standards means every KPI in the P&P, Handbook, store scorecard and signed updates on the 1st.",
    "Miss one and that month's bonus isn't earned — for the store, or just you if the miss is yours.",
    "Management approves every bonus and can suspend a store's program with written notice.",
]
RATE_NOTE = None
LEGAL = ("I received, understand, and will follow this policy. The company may change or end it at any time "
         "without notice. This is not an employment contract and does not change at-will employment; either "
         "the company or I may end my employment at any time, with or without cause or notice.")

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(0.5)
s.left_margin = s.right_margin = Inches(0.75)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(12)


def para(text, size, bold=False, color=None, after=0, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return p


para("VALLEY PAWN", 9, bold=True, color=BLUE, after=0)
para(TITLE, 22, bold=True, color=NAVY, after=0)
para(STARTS, 9.5, color=GREY, after=10)

box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.rows[0].cells[0]
cell.width = Inches(7.0)
sh = OxmlElement("w:shd")
sh.set(qn("w:val"), "clear")
sh.set(qn("w:fill"), "2D1A5E")
cell._tc.get_or_add_tcPr().append(sh)
cell.text = ""
for i, line in enumerate(BOX):
    p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10 if i == 0 else 0)
    p.paragraph_format.space_after = Pt(10 if i == len(BOX) - 1 else 2)
    r = p.add_run(line)
    r.font.size = Pt(17)
    r.bold = True
    r.font.color.rgb = WHITE

para("", 6, after=4)
for line in LINES:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    r = p.add_run("→  ")
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = BLUE
    r2 = p.add_run(line)
    r2.font.size = Pt(12)

if RATE_NOTE:
    para(RATE_NOTE, 9.5, color=GREY, before=2, after=0)
para(LEGAL, 8, color=GREY, before=10, after=6)

doc.save(DOCX)
print("wrote", DOCX)
