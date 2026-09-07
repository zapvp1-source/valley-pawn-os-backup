#!/usr/bin/env python3
"""
bonus_engine.py — deterministic Valley Pawn bonus engine (Full Circle Finance Inc).

Every number and every Slack body for the bonus program comes out of this file.
Rules live in ../bonus_rules.json. Inputs are pipeline files only — never Bravo GUI, never QBO.

Stages
  collect --month YYYY-MM   copy + validate the month's source files into data/YYYY-MM/
  close   --month YYYY-MM   qualifiers + payouts for an earning month -> out/YYYY-MM/, ledger
  targets --month YYYY-MM   next-month targets from the completed month (trailing-12 yield)

Exit codes: 0 = complete, 2 = HELD (a gate failed; nothing may be posted), 1 = usage/crash.
Rule 18: on exit 2 the Slack bodies are NOT written (out/ only gets hold.txt).
"""
import argparse, csv, json, os, re, shutil, sys, datetime as dt
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    print("openpyxl missing", file=sys.stderr); sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT_DEFAULT = os.path.abspath(os.path.join(HERE, '..'))
STORES = ['CUL', 'HAR', 'LEX', 'ROA', 'WAY']
MONTHS = ['January','February','March','April','May','June','July','August','September','October','November','December']

# ----------------------------------------------------------------------------- helpers

def money(v):
    if v is None: return None
    s = str(v).strip().replace('$', '').replace(',', '')
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()')
    try:
        return -float(s) if neg else float(s)
    except Exception:
        return None

def fmt(v):
    return '—' if v is None else ('$' + format(round(v, 2), ',.2f'))

def fmt0(v):
    return '—' if v is None else ('$' + format(round(v), ',.0f'))

def last_day(month):
    y, m = map(int, month.split('-'))
    nxt = dt.date(y + (m == 12), (m % 12) + 1, 1)
    return nxt - dt.timedelta(days=1)

def prev_month(month):
    y, m = map(int, month.split('-'))
    return f"{y-1}-12" if m == 1 else f"{y}-{m-1:02d}"

def next_month(month):
    y, m = map(int, month.split('-'))
    return f"{y+1}-01" if m == 12 else f"{y}-{m+1:02d}"

def month_label(month):
    y, m = map(int, month.split('-'))
    return f"{MONTHS[m-1]} {y}"

def payday(month):
    """First Friday after the 15th of the month AFTER the earning month."""
    nm = next_month(month); y, m = map(int, nm.split('-'))
    d = dt.date(y, m, 16)
    while d.weekday() != 4: d += dt.timedelta(days=1)
    return d

class Ctx:
    def __init__(self, root, pipeline, month):
        self.root = root; self.pipeline = pipeline; self.month = month
        self.rules = json.load(open(os.path.join(root, 'bonus_rules.json')))
        self.data = os.path.join(root, 'data', month)
        self.out = os.path.join(root, 'out', month)
        os.makedirs(self.data, exist_ok=True); os.makedirs(self.out, exist_ok=True)
        self.gaps = []; self.notes = []; self.holds = []
    def regime(self, month=None):
        month = month or self.month
        for r in self.rules['regimes']:
            if r['effective_from'] <= month: return r
        return self.rules['regimes'][-1]

# ----------------------------------------------------------------------------- EOM parsing

def eom_reporting_dates(path):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).active
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        vals = [v for v in row if v is not None]
        if vals and str(vals[0]).startswith('Reporting Dates') and len(vals) > 1:
            return str(vals[1]).strip()
    return None

def eom_is_month(path, month):
    rd = eom_reporting_dates(path)
    if not rd: return False, rd
    ld = last_day(month)
    want = f"{int(month[5:7])}/1/{month[:4]} - {ld.month}/{ld.day}/{ld.year}"
    return rd == want, rd

def parse_eom(path):
    """Verified formula (bonus_kpis_extract.py / store_kpis_compile.py):
    PSC = In-Store Subtotal Interest+Fees+Misc; Net Revenue = PSC + Sales Revenue (Profit)."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    MC = ws.max_column
    def nn(r): return [x for x in [ws.cell(r, c).value for c in range(1, MC + 1)] if x is not None]
    def find(lbl):
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v is not None and str(v).strip().startswith(lbl): return r
        return 0
    L = nn(find('Ending Loan Base')); I = nn(find('Ending Inventory Base'))
    loan = money(L[2]) if len(L) > 2 else 0.0
    inv = money(I[2]) if len(I) > 2 else 0.0
    sub = nn(find('In-Store Subtotal'))
    psc = sum((money(sub[i]) or 0.0) for i in (4, 5, 6) if len(sub) > i)
    rev = nn(find('Sales Revenue (Profit)'))
    prof = (money(rev[-1]) or 0.0) if rev else 0.0
    ist = nn(find('In-Store Total')); st = nn(find('Sales Total'))
    txn = (money(ist[1]) or 0) + (money(sub[1]) or 0) + (money(st[1]) or 0)
    return {'net_revenue': round(psc + prof, 2), 'psc': round(psc, 2), 'sales_profit': round(prof, 2),
            'loan_balance': round(loan or 0, 2), 'inventory_balance': round(inv or 0, 2),
            'ending_assets': round((loan or 0) + (inv or 0), 2), 'txn_volume': int(txn)}

def find_month_eom(ctx, month, store):
    """Locate a month-only EOM file for store/month: sidecar first, then output/, validated by Reporting Dates."""
    ld = last_day(month).isoformat()
    y, m = month.split('-')
    cands = [os.path.join(ctx.pipeline, 'output', 'monthly-analytics', month, f'same-month-current_{store}.xlsx'),
             # the prior-year copy of a month lives in the FOLLOWING year's sidecar as same-month-prior
             os.path.join(ctx.pipeline, 'output', 'monthly-analytics', f'{int(y)+1}-{m}', f'same-month-prior_{store}.xlsx'),
             os.path.join(ctx.pipeline, 'output', f'{ld}_{store}_end-of-month.xlsx')]
    # bonus-owned isolated pulls
    cands.insert(0, os.path.join(ctx.pipeline, 'output', 'bonus', month, f'{ld}_{store}_end-of-month.xlsx'))
    for p in cands:
        if os.path.exists(p) and os.path.getsize(p) > 500:
            ok, rd = eom_is_month(p, month)
            if ok: return p, rd
    return None, None

# ----------------------------------------------------------------------------- collect

def stage_collect(ctx):
    rep = {'month': ctx.month, 'eom': {}, 'chekkit': {}, 'employee_activity': {}, 'gold': {}}
    ld = last_day(ctx.month).isoformat()
    for s in STORES:
        p, rd = find_month_eom(ctx, ctx.month, s)
        if p:
            dst = os.path.join(ctx.data, f'eom_{s}.xlsx'); shutil.copy2(p, dst)
            rep['eom'][s] = {'src': p, 'reporting_dates': rd}
        else:
            rep['eom'][s] = None; ctx.holds.append(f'EOM month file missing/invalid for {s}')
        src = os.path.join(ctx.pipeline, 'output', f'{ld}_{s}_chekkit-invites-range.csv')
        if os.path.exists(src):
            shutil.copy2(src, os.path.join(ctx.data, f'chekkit_invites_{s}.csv')); rep['chekkit'][s] = src
        else:
            rep['chekkit'][s] = None; ctx.gaps.append(f'Chekkit invites CSV missing for {s} ({ld})')
        src = os.path.join(ctx.pipeline, 'output', f'{ld}_{s}_employee-activity-range.csv')
        if os.path.exists(src):
            shutil.copy2(src, os.path.join(ctx.data, f'employee_activity_{s}.csv')); rep['employee_activity'][s] = src
        else:
            rep['employee_activity'][s] = None; ctx.gaps.append(f'Employee activity (month range) CSV missing for {s}')
        src = os.path.join(ctx.pipeline, 'output', f'{ctx.month[:4]}_{s}_scrap-refining-gold.csv')
        if os.path.exists(src):
            shutil.copy2(src, os.path.join(ctx.data, f'gold_{s}.csv')); rep['gold'][s] = src
        else:
            rep['gold'][s] = None; ctx.gaps.append(f'Scrap gold CSV missing for {s}')
    rep['gaps'] = ctx.gaps; rep['holds'] = ctx.holds
    json.dump(rep, open(os.path.join(ctx.data, 'collect_report.json'), 'w'), indent=1)
    print(json.dumps(rep, indent=1))
    return 2 if ctx.holds else 0

# ----------------------------------------------------------------------------- readers

def read_json(ctx, name, required=False):
    p = os.path.join(ctx.data, name)
    if os.path.exists(p): return json.load(open(p))
    if required: ctx.holds.append(f'{name} missing in data/{ctx.month}')
    else: ctx.gaps.append(f'{name} missing in data/{ctx.month}')
    return None

def read_email_pct(ctx, s):
    p = os.path.join(ctx.data, f'chekkit_invites_{s}.csv')
    if not os.path.exists(p): return None, 0, 0
    rows = list(csv.DictReader(open(p, newline='', encoding='utf-8-sig')))
    tot = len(rows)
    cap = sum(1 for r in rows if (r.get('email') or '').strip() and '@' in r.get('email', ''))
    return (round(cap / tot, 4) if tot else None), cap, tot

def read_gold(ctx, s):
    p = os.path.join(ctx.data, f'gold_{s}.csv')
    if not os.path.exists(p): return None, []
    best = {}
    for r in csv.DictReader(open(p, newline='', encoding='utf-8-sig')):
        key = (r['BucketName'].strip(), r['CreatedOn'].strip())
        cur = best.get(key)
        if cur is None or (r['Status'] == 'CLOSED' and cur['Status'] != 'CLOSED'): best[key] = r
    tot = 0.0; used = []
    for r in best.values():
        if r['Status'] != 'CLOSED': continue
        m = re.match(r'(\d+)/(\d+)/(\d{4})', r['StatusDate'].strip())
        if not m: continue
        ym = f"{m.group(3)}-{int(m.group(1)):02d}"
        if ym == ctx.month:
            w = float(r['CombinedMetalWeightDwt'] or 0); tot += w
            used.append({'bucket': r['BucketName'], 'closed': r['StatusDate'], 'dwt': round(w, 2)})
    return round(tot, 2), used

def read_employee_activity(ctx, s):
    p = os.path.join(ctx.data, f'employee_activity_{s}.csv')
    if not os.path.exists(p): return None, None
    rows = list(csv.reader(open(p, newline='', encoding='utf-8-sig', errors='replace')))
    rd = None; hdr_i = None
    for i, r in enumerate(rows):
        if len(r) > 2 and str(r[2]).startswith('Reporting Dates') and len(r) > 9: rd = r[9]
        if r and r[0] == 'Employee': hdr_i = i; break
    if hdr_i is None: return rd, None
    hdr = rows[hdr_i]; gi = hdr.index('Retail Gross Profit'); si = hdr.index('Retail Sales Excluding Fees')
    out = {}
    for r in rows[hdr_i + 1:]:
        if not r or not r[0] or r[0].startswith('Report printed') or r[0] == 'Total Store': continue
        code = r[0].split(' - ')[0].strip(); name = r[0].split(' - ', 1)[-1].strip()
        out[code] = {'name': name, 'retail_gp': money(r[gi]) or 0.0, 'retail_sales': money(r[si]) or 0.0}
    return rd, out

def norm(s): return re.sub(r'[^a-z]', '', (s or '').lower())

def map_bravo_to_roster(ctx, code, bname, roster):
    """Match a Bravo employee row to a Gusto roster entry by name."""
    parts = bname.replace(',', ' ').split()
    if not parts: return None
    first, last = norm(parts[0]), norm(parts[-1])
    emap = ctx.rules.get('_employee_map', {})
    if code in emap:
        for e in roster:
            if e['uuid'] == emap[code]: return e
    for e in roster:
        el = norm(e['last_name']); ef = norm(e['first_name']); ep = norm(e.get('preferred_first_name') or '')
        if el == last and (ef == first or ep == first or ef.startswith(first) or first.startswith(ef)): return e
    for e in roster:
        if norm(e['last_name']) == last: return e
    return None

# ----------------------------------------------------------------------------- close

def stage_close(ctx):
    R = ctx.rules; reg = ctx.regime(); month = ctx.month; ld = last_day(month)
    tq = reg['task_qualifiers']
    targets = read_json(ctx, 'targets.json', required=True) or {}
    roster = read_json(ctx, 'roster.json', required=True) or []
    reviews = read_json(ctx, 'reviews.json') or {}
    fb = read_json(ctx, 'fb_gains.json') or {}
    history = load_history(ctx)  # for plausibility band

    stores = {}
    for s in STORES:
        p = os.path.join(ctx.data, f'eom_{s}.xlsx')
        if not os.path.exists(p): ctx.holds.append(f'eom_{s}.xlsx missing — run collect'); continue
        ok, rd = eom_is_month(p, month)
        if not ok: ctx.holds.append(f'eom_{s}.xlsx Reporting Dates "{rd}" is not {month}'); continue
        k = parse_eom(p)
        # plausibility vs trailing-12 average
        hist = [history.get(m, {}).get(s, {}).get('net_revenue') for m in trailing_months(prev_month(month), 12)]
        hist = [h for h in hist if h]
        if len(hist) >= 6:
            avg = sum(hist) / len(hist); band = R['gates']['net_revenue_plausibility_band']
            if not (avg * (1 - band) <= k['net_revenue'] <= avg * (1 + band)):
                ctx.holds.append(f'{s} net revenue {fmt(k["net_revenue"])} outside ±{int(band*100)}% of trailing-12 avg {fmt(avg)} — wrong file?')
        tgt = (targets.get('targets') or {}).get(s); py = (targets.get('prior_year') or {}).get(s)
        if tgt is None: ctx.holds.append(f'no target for {s} {month}')
        em, cap, tot = read_email_pct(ctx, s)
        gold, buckets = read_gold(ctx, s)
        rv = reviews.get(s) if isinstance(reviews, dict) else None
        fg = fb.get(s) if isinstance(fb, dict) else None
        st = dict(k, target=tgt, prior_year=py, email_pct=em, email_cap=cap, email_tot=tot, gold_dwt=gold,
                  gold_buckets=buckets, reviews=rv, fb_gain=fg)
        st['bridge1'] = (tgt is not None and k['net_revenue'] >= tgt)
        st['bridge2'] = (py is not None and k['net_revenue'] >= py)
        q = {}
        q['reviews'] = None if rv is None else rv >= tq['reviews_min']
        q['email'] = None if em is None else em >= tq['email_pct_min']
        q['gold'] = None if gold is None else gold >= tq['gold_dwt_min']
        if tq.get('fb_follower_gain_min') is not None:
            q['fb'] = None if fg is None else fg >= tq['fb_follower_gain_min']
        if tq.get('revenue_yoy_required'):
            q['rev_yoy'] = None if py is None else st['bridge2']
        if tq.get('social_qr_views_required'):
            q['social_qr'] = None  # rail retired with the Aug-2026 regime; July closes are HELD on this unless gold/reviews/email already decide it
        st['qualifiers'] = q
        vals = list(q.values())
        if any(v is False for v in vals): st['tier2'] = False
        elif all(v is True for v in vals): st['tier2'] = True
        else: st['tier2'] = None  # HELD
        stores[s] = st

    if ctx.holds:
        return hold(ctx)

    # category wins + top store
    cats = {'email_pct': 'Email %', 'gold_dwt': 'Gold dwt', 'reviews': 'Reviews'}
    if tq.get('fb_follower_gain_min') is not None: cats['fb_gain'] = 'FB gain'
    wins = defaultdict(list)
    for c in cats:
        vals = {s: stores[s].get(c) for s in STORES if stores[s].get(c) is not None}
        if not vals: continue
        top = max(vals.values())
        for s, v in vals.items():
            if v == top: wins[s].append(cats[c])
    top_store = None
    if wins:
        best = max(len(v) for v in wins.values())
        tied = [s for s in STORES if len(wins.get(s, [])) == best]
        if len(tied) > 1: tied.sort(key=lambda s: -(stores[s].get('gold_dwt') or 0))
        top_store = tied[0]

    # payouts
    pay = []; preston_count = sum(1 for s in STORES if stores[s]['bridge1'])
    active = [e for e in roster if not e.get('terminated')]
    for s in STORES:
        st = stores[s]; name = R['stores'][s]
        rd, act = read_employee_activity(ctx, s)
        if act is None: ctx.gaps.append(f'{name}: employee activity file missing — associate payouts not computed')
        if rd and not rd.startswith(f"{int(month[5:7])}/1/{month[:4]}"):
            ctx.gaps.append(f'{name}: employee activity range is "{rd}", expected {month}')
        if not st['bridge1']:
            pay.append({'store': s, 'employee': '(store)', 'role': '—', 'basis': None, 'rate': 0, 'bonus': 0.0,
                        'note': f'Bridge 1 not met ({fmt(st["net_revenue"])} vs target {fmt(st["target"])}) — $0 for the store'})
            continue
        tier2 = st['tier2']; held = tier2 is None
        mrate = reg['rates']['manager']['hit' if tier2 else 'miss']
        arate = reg['rates']['associate']['hit' if tier2 else 'miss']
        # manager
        mgrs = [e for e in active if e['department'] == name and e['title'] in R['eligibility']['manager_titles']]
        elig_mgrs = []
        for e in mgrs:
            hd = dt.date.fromisoformat(e['hire_date'])
            if R['eligibility']['manager_full_month'] and hd > dt.date(int(month[:4]), int(month[5:7]), 1):
                ctx.notes.append(f'{name}: {e["display"]} is titled Manager but hired {hd} — not a full month, treated as associate-tier for {month}')
            else: elig_mgrs.append(e)
        if not elig_mgrs:
            # de facto lead = highest retail sales associate in Bravo at that store
            leads = []
            if act:
                for code, a in act.items():
                    if code in R['bravo_ignore_codes']: continue
                    e = map_bravo_to_roster(ctx, code, a['name'], active)
                    if e and e['department'] == name and e['title'] in R['eligibility']['associate_titles']:
                        leads.append((a['retail_sales'], e))
            if leads:
                leads.sort(key=lambda t: -t[0]); lead = leads[0][1]; elig_mgrs = [lead]
                ctx.notes.append(f'{name}: no titled Manager for the full month — {lead["display"]} paid on the Manager formula as de facto lead (flagged)')
            else:
                ctx.gaps.append(f'{name}: no Manager and no de facto lead found — manager-tier payout not computed')
        for e in elig_mgrs:
            b = st['net_revenue'] * mrate
            pay.append({'store': s, 'employee': e['display'], 'uuid': e['uuid'], 'role': 'Manager' + ('' if e['title'] == 'Manager' else ' (de facto)'),
                        'basis': st['net_revenue'], 'rate': mrate, 'bonus': round(b, 2), 'held': held,
                        'note': 'Tier 2 ' + ('HELD' if held else ('hit' if tier2 else 'missed'))})
        # associates from Bravo per-employee retail GP
        if act:
            for code, a in act.items():
                if code in R['bravo_ignore_codes']: continue
                e = map_bravo_to_roster(ctx, code, a['name'], roster)
                if e is None:
                    if a['retail_gp'] > 0: ctx.gaps.append(f'{name}: Bravo user {code} ({a["name"]}) not in Gusto roster — GP {fmt(a["retail_gp"])} unpaid')
                    continue
                if e.get('terminated'):
                    if a['retail_gp'] > 0: ctx.notes.append(f'{name}: {e["display"]} earned GP {fmt(a["retail_gp"])} but is no longer employed — ineligible at payout')
                    continue
                if e in elig_mgrs or e['title'] not in R['eligibility']['associate_titles']: continue
                if a['retail_gp'] <= 0: continue
                b = a['retail_gp'] * arate
                pay.append({'store': s, 'employee': e['display'], 'uuid': e['uuid'], 'role': e['title'],
                            'basis': a['retail_gp'], 'rate': arate, 'bonus': round(b, 2), 'held': held,
                            'note': ('home store ' + e['department'] + '; ' if e['department'] != name else '') + 'Tier 2 ' + ('HELD' if held else ('hit' if tier2 else 'missed'))})
    # top store pool
    if top_store and stores[top_store]['bridge1']:
        home = R['stores'][top_store]
        members = [p for p in pay if p['store'] == top_store and p.get('uuid')
                   and next((e for e in active if e['uuid'] == p['uuid']), {}).get('department') == home]
        if members:
            share = round(reg['top_store_pool'] / len(members), 2)
            for p in members: p['top_store_share'] = share
            pay.append({'store': top_store, 'employee': '(Top store pool)', 'role': 'pool', 'basis': None, 'rate': None,
                        'bonus': reg['top_store_pool'], 'note': f'{R["stores"][top_store]} — {len(wins[top_store])} category wins: {", ".join(wins[top_store])}; split {fmt(share)} x {len(members)}'})
    elif top_store:
        ctx.notes.append(f'Top performer {R["stores"][top_store]} did not pass Bridge 1 — $300 pool not paid')
    preston = next((e for e in active if e['title'] == 'Market Manager'), None)
    if preston:
        pay.append({'store': 'ALL', 'employee': preston['display'], 'uuid': preston['uuid'], 'role': 'Market Manager',
                    'basis': preston_count, 'rate': reg['market_manager_per_bridge1_store'],
                    'bonus': round(preston_count * reg['market_manager_per_bridge1_store'], 2),
                    'note': f'{preston_count}/5 stores hit target x {fmt0(reg["market_manager_per_bridge1_store"])}'})
    total = round(sum(p['bonus'] for p in pay if p.get('uuid')) , 2)
    held_any = any(p.get('held') for p in pay)

    result = {'month': month, 'regime': reg['name'], 'payday': payday(month).isoformat(), 'stores': stores,
              'category_wins': dict(wins), 'top_store': top_store, 'payouts': pay, 'total_payout': total,
              'held': held_any, 'gaps': ctx.gaps, 'notes': ctx.notes, 'field_posting': R['field_posting']['enabled'],
              'generated': dt.datetime.now().isoformat(timespec='seconds')}
    json.dump(result, open(os.path.join(ctx.out, 'close.json'), 'w'), indent=1, default=str)
    write_close_outputs(ctx, result)
    update_ledger(ctx, result)
    print(open(os.path.join(ctx.out, 'close_report.md')).read())
    return 0

def hold(ctx):
    msg = '\n'.join(ctx.holds)
    open(os.path.join(ctx.out, 'hold.txt'), 'w').write(f'HELD {ctx.month} {dt.datetime.now():%Y-%m-%d %H:%M}\n{msg}\n')
    print(f'HELD — nothing may be posted for {ctx.month}:\n{msg}', file=sys.stderr)
    return 2

def trailing_months(end_month, n):
    out = []; m = end_month
    for _ in range(n): out.append(m); m = prev_month(m)
    return list(reversed(out))

def load_history(ctx):
    """month -> store -> {net_revenue, ending_assets} from every validated month-only EOM file we can find (cached)."""
    cache_p = os.path.join(ctx.root, 'data', 'history.json')
    hist = json.load(open(cache_p)) if os.path.exists(cache_p) else {}
    changed = False
    start = dt.date(2025, 1, 1); m = f"{start.year}-{start.month:02d}"
    while m <= ctx.month:
        for s in STORES:
            if hist.get(m, {}).get(s): continue
            p, rd = find_month_eom(ctx, m, s)
            if p:
                k = parse_eom(p); hist.setdefault(m, {})[s] = {'net_revenue': k['net_revenue'], 'ending_assets': k['ending_assets'], 'src': p}
                changed = True
        m = next_month(m)
    if changed:
        os.makedirs(os.path.dirname(cache_p), exist_ok=True)
        json.dump(hist, open(cache_p, 'w'), indent=1)
    return hist

# ----------------------------------------------------------------------------- outputs

def write_close_outputs(ctx, res):
    R = ctx.rules; S = R['stores']; st = res['stores']; ml = month_label(ctx.month)
    def yn(v): return '✅' if v else ('❌' if v is False else '—')
    # 1. qualifiers table (field-safe: no bridge/tier words)
    lines = [f"🎯 *Bonus Qualifiers — {ml}*", "", "```",
             f"{'Store':<13}{'Revenue':>11}{'Target':>11} Hit  {'Reviews':>8} {'Email':>7} {'Gold':>8} {'FB':>5}  Task",
             '-' * 78]
    for s in STORES:
        x = st[s]
        lines.append(f"{S[s]:<13}{fmt0(x['net_revenue']):>11}{fmt0(x['target']):>11}  {yn(x['bridge1'])}  "
                     f"{(x['reviews'] if x['reviews'] is not None else '—'):>8} "
                     f"{(format(x['email_pct']*100, '.0f') + '%') if x['email_pct'] is not None else '—':>7} "
                     f"{(format(x['gold_dwt'], '.1f') if x['gold_dwt'] is not None else '—'):>8} "
                     f"{(('+' if (x['fb_gain'] or 0) > 0 else '') + str(x['fb_gain'])) if x['fb_gain'] is not None else '—':>5}  "
                     f"{yn(x['tier2']) if x['bridge1'] else '—'}")
    lines += ['```', f"🏆 Top performer: {S[res['top_store']] if res['top_store'] else '—'}"]
    open(os.path.join(ctx.out, 'slack_qualifiers.txt'), 'w').write('\n'.join(lines) + '\n')
    # 2. Joshua DM (dollar figures OK)
    d = [f"*Bonus Payouts — {ml}* ({res['regime']}) — payday {res['payday']}", ""]
    for s in STORES:
        x = st[s]; rows = [p for p in res['payouts'] if p['store'] == s]
        head = f"*{S[s]}* — {fmt(x['net_revenue'])} vs target {fmt(x['target'])} → {'HIT' if x['bridge1'] else 'MISSED'}"
        if x['bridge1']:
            head += f"; YoY {'up' if x['bridge2'] else 'down'} (2025: {fmt(x['prior_year'])}); task bonus {'HIT' if x['tier2'] else ('HELD' if x['tier2'] is None else 'missed')}"
            q = x['qualifiers']
            head += f" [reviews {x['reviews'] if x['reviews'] is not None else '—'} {yn(q.get('reviews'))}, email {format((x['email_pct'] or 0)*100,'.0f')}% {yn(q.get('email'))}, gold {x['gold_dwt']} {yn(q.get('gold'))}" + (f", FB {x['fb_gain']:+d} {yn(q.get('fb'))}" if 'fb' in q else '') + "]"
        d.append(head)
        for p in rows:
            if p.get('uuid'):
                d.append(f"  • {p['employee']} ({p['role']}): {fmt(p['basis'])} × {p['rate']*100:.1f}% = *{fmt(p['bonus'])}*" + (f" + {fmt(p['top_store_share'])} top-store share" if p.get('top_store_share') else '') + (' ⚠️ HELD' if p.get('held') else ''))
            elif p['role'] == 'pool': d.append(f"  • Top-store pool {fmt(p['bonus'])}: {p['note']}")
            else: d.append(f"  • {p['note']}")
        d.append('')
    pr = next((p for p in res['payouts'] if p['role'] == 'Market Manager'), None)
    if pr: d.append(f"*{pr['employee']}* (Market Manager): {pr['note']} = *{fmt(pr['bonus'])}*")
    d.append(f"\n*TOTAL {fmt(res['total_payout'])}*" + (' — contains HELD lines, not final' if res['held'] else ''))
    if res['notes']: d += ['', '_Notes_'] + [f"• {n}" for n in res['notes']]
    if res['gaps']: d += ['', '_Gaps (need attention)_'] + [f"• {g}" for g in res['gaps']]
    d.append("\nReply *approve* and the lines get loaded into the payday payroll for you to submit; *hold* keeps everything parked.")
    open(os.path.join(ctx.out, 'dm_payout.txt'), 'w').write('\n'.join(d) + '\n')
    # 3. Gusto lines
    with open(os.path.join(ctx.out, 'gusto_lines.csv'), 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['employee_uuid', 'employee', 'store', 'bonus', 'memo', 'held'])
        for p in res['payouts']:
            if p.get('uuid'):
                amt = round(p['bonus'] + p.get('top_store_share', 0), 2)
                w.writerow([p['uuid'], p['employee'], p['store'], f"{amt:.2f}", f"{ml} bonus — {p['role']}", 'Y' if p.get('held') else 'N'])
    # 4. report
    r = [f"# Bonus close — {ml}", f"Generated {res['generated']} · regime: {res['regime']} · payday {res['payday']} · field posting {'ON' if res['field_posting'] else 'OFF'}", "",
         "## Stores", "| Store | Net rev | Target | Hit | 2025 | YoY | Reviews | Email | Gold | FB | Task | Assets | Txns |", "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for s in STORES:
        x = st[s]
        r.append(f"| {S[s]} | {fmt(x['net_revenue'])} | {fmt(x['target'])} | {yn(x['bridge1'])} | {fmt(x['prior_year'])} | {yn(x['bridge2'])} | {x['reviews']} | {x['email_pct']} ({x['email_cap']}/{x['email_tot']}) | {x['gold_dwt']} | {x['fb_gain']} | {yn(x['tier2'])} | {fmt(x['ending_assets'])} | {x['txn_volume']} |")
    r += ["", f"Category wins: {res['category_wins']} → top store {res['top_store']}", "", "## Payouts", "| Store | Employee | Role | Basis | Rate | Bonus | Note |", "|---|---|---|---|---|---|---|"]
    for p in res['payouts']:
        r.append(f"| {p['store']} | {p['employee']} | {p['role']} | {fmt(p['basis']) if isinstance(p['basis'], float) else p['basis']} | {p['rate']} | {fmt(p['bonus'])} | {p['note']} |")
    r += ["", f"**Total {fmt(res['total_payout'])}**" + (' (HELD lines present)' if res['held'] else '')]
    if res['notes']: r += ["", "## Notes"] + [f"- {n}" for n in res['notes']]
    if res['gaps']: r += ["", "## Gaps"] + [f"- {g}" for g in res['gaps']]
    r += ["", "## Gold buckets used"] + [f"- {S[s]}: " + ('; '.join(f"{b['bucket']} closed {b['closed'][:10]} {b['dwt']} dwt" for b in st[s]['gold_buckets']) or 'none') for s in STORES]
    open(os.path.join(ctx.out, 'close_report.md'), 'w').write('\n'.join(r) + '\n')

# ----------------------------------------------------------------------------- ledger

def update_ledger(ctx, res):
    p = os.path.join(ctx.root, 'ledger', 'VP_Bonus_Ledger_2026.xlsx'); os.makedirs(os.path.dirname(p), exist_ok=True)
    wb = openpyxl.load_workbook(p) if os.path.exists(p) else openpyxl.Workbook()
    if 'Sheet' in wb.sheetnames and len(wb.sheetnames) == 1: wb.remove(wb['Sheet'])
    S = ctx.rules['stores']
    def sheet(name, header):
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name); ws.append(header)
        return wb[name]
    q = sheet('Qualifiers', ['Month', 'Store', 'Net Revenue', 'Target', 'Target Hit', '2025 Revenue', 'YoY Up', 'Reviews', 'Email %', 'Email Captured', 'Email Opps', 'Gold dwt', 'FB Gain', 'Task Bonus', 'Ending Assets', 'Txn Volume', 'Category Wins', 'Top Store', 'Regime', 'Generated'])
    py = sheet('Payouts', ['Month', 'Store', 'Employee', 'Role', 'Basis', 'Rate', 'Bonus', 'Top-store share', 'Held', 'Note', 'Gusto UUID', 'Payday'])
    tr = sheet('Trend', ['Month', 'Company Net Revenue', 'Stores Hit', 'Total Payout', 'Held', 'Top Store'])
    for ws, col in ((q, 1), (py, 1), (tr, 1)):
        for r in range(ws.max_row, 1, -1):
            if ws.cell(r, col).value == ctx.month: ws.delete_rows(r)
    st = res['stores']
    for s in STORES:
        x = st[s]
        q.append([ctx.month, S[s], x['net_revenue'], x['target'], 'Y' if x['bridge1'] else 'N', x['prior_year'], 'Y' if x['bridge2'] else 'N', x['reviews'], x['email_pct'], x['email_cap'], x['email_tot'], x['gold_dwt'], x['fb_gain'], {True: 'Y', False: 'N', None: 'HELD'}[x['tier2']] if x['bridge1'] else '—', x['ending_assets'], x['txn_volume'], ', '.join(res['category_wins'].get(s, [])), 'Y' if res['top_store'] == s else '', res['regime'], res['generated']])
    for p_ in res['payouts']:
        py.append([ctx.month, p_['store'], p_['employee'], p_['role'], p_['basis'] if isinstance(p_['basis'], (int, float)) else None, p_['rate'], p_['bonus'], p_.get('top_store_share'), 'Y' if p_.get('held') else '', p_['note'], p_.get('uuid'), res['payday']])
    tr.append([ctx.month, round(sum(st[s]['net_revenue'] for s in STORES), 2), sum(1 for s in STORES if st[s]['bridge1']), res['total_payout'], 'Y' if res['held'] else '', S.get(res['top_store'], '')])
    for ws in (q, py, tr):
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(10, min(48, max(len(str(c.value or '')) for c in col) + 2))
    wb.save(p)

# ----------------------------------------------------------------------------- targets

def stage_targets(ctx):
    """ctx.month = completed month. Produces targets for next month."""
    R = ctx.rules; S = R['stores']; hist = load_history(ctx); tm = next_month(ctx.month)
    window = trailing_months(ctx.month, 12)
    res = {'completed_month': ctx.month, 'target_month': tm, 'method': R['targets_method']['name'], 'stores': {}, 'targets': {}, 'prior_year': {}}
    for s in STORES:
        if not hist.get(ctx.month, {}).get(s): ctx.holds.append(f'{s}: no validated month-only EOM for {ctx.month}')
    if ctx.holds: return hold(ctx)
    for s in STORES:
        ys = []
        for m in window:
            cur = hist.get(m, {}).get(s); prv = hist.get(prev_month(m), {}).get(s)
            if cur and prv and prv['ending_assets']: ys.append(cur['net_revenue'] / prv['ending_assets'])
        if len(ys) < 12: ctx.gaps.append(f'{s}: only {len(ys)}/12 monthly yields available in the trailing window')
        if len(ys) < 9: ctx.holds.append(f'{s}: fewer than 9 yields — target not defensible')
        y = sum(ys) / len(ys) if ys else None
        ea = hist[ctx.month][s]['ending_assets']
        tgt = round(ea * y) if y else None
        res['stores'][s] = {'ending_assets': ea, 'trail12_yield': round(y, 4) if y else None, 'n_yields': len(ys), 'target': tgt}
        res['targets'][s] = tgt
        pym = f"{int(tm[:4])-1}-{tm[5:]}"
        res['prior_year'][s] = hist.get(pym, {}).get(s, {}).get('net_revenue')
    if ctx.holds: return hold(ctx)
    res['company_target'] = sum(res['targets'].values()); res['gaps'] = ctx.gaps
    json.dump(res, open(os.path.join(ctx.out, 'targets.json'), 'w'), indent=1)
    # seed next month's data folder so close can find the targets
    nd = os.path.join(ctx.root, 'data', tm); os.makedirs(nd, exist_ok=True)
    json.dump({'targets': res['targets'], 'prior_year': res['prior_year'], 'source': f'bonus_engine targets from {ctx.month}', 'generated': dt.datetime.now().isoformat(timespec='seconds')}, open(os.path.join(nd, 'targets.json'), 'w'), indent=1)
    # Slack body — Joshua's approved template
    ml = MONTHS[int(tm[5:7]) - 1]; pm = MONTHS[int(ctx.month[5:7]) - 1]; y0 = window[0]
    L = [f"📅 {ml} {tm[:4]} Bonus Targets", f"{ml} Targets by Store"]
    for s in ['CUL', 'HAR', 'ROA', 'LEX', 'WAY']: L.append(f"🏪 {S[s]} — {fmt0(res['targets'][s])}")
    L += ["", "How We Got Here.",
          f"-Targets are built using each store's own trailing 12-month yield — not a company average. Formula: {pm} ending assets × store's trailing-12-month yield.",
          f"-{pm} Ending Assets (Loans + Inventory):",
          ' · '.join(f"{S[s]} {fmt0(res['stores'][s]['ending_assets'])}" for s in ['CUL', 'HAR', 'ROA', 'LEX', 'WAY']),
          f"-Trail-12 Avg Yield ({MONTHS[int(y0[5:7])-1][:3]} {y0[:4]}–{pm[:3]} {ctx.month[:4]}): " + ' · '.join(f"{S[s]} {res['stores'][s]['trail12_yield']*100:.1f}%" for s in ['CUL', 'HAR', 'ROA', 'LEX', 'WAY']),
          f"Each store's number reflects what they've actually been doing this year. Hit it and earn it. This not math anymore, it's science. Let's have a great {ml}. 💪"]
    open(os.path.join(ctx.out, 'slack_targets.txt'), 'w').write('\n'.join(L) + '\n')
    print(json.dumps(res, indent=1)); print('\n'.join(L))
    return 0

# ----------------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('stage', choices=['collect', 'close', 'targets'])
    ap.add_argument('--month', required=True, help='YYYY-MM (earning month for collect/close; completed month for targets)')
    ap.add_argument('--root', default=ROOT_DEFAULT)
    ap.add_argument('--pipeline', default=None, help='Bravo Data Extraction folder')
    a = ap.parse_args()
    pipeline = a.pipeline or os.path.join(os.path.dirname(a.root), 'Bravo Data Extraction')
    ctx = Ctx(a.root, pipeline, a.month)
    rc = {'collect': stage_collect, 'close': stage_close, 'targets': stage_targets}[a.stage](ctx)
    sys.exit(rc)

if __name__ == '__main__':
    main()
