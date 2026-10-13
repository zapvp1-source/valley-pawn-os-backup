import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as L
d=json.load(open('out10.json')); T=[x for x in d['tickets'] if x['paid']]
sp=d['spot']
F=Font(name='Arial',size=10); H=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F3864')
cols=[('Store','store',7),('Ticket','ticket',14),('Customer','customer',24),('Phone','phone',15),('Email','email',24),('SMS','sms',8),('Description','desc',48),
('Loan $','loan',10),('Loan on metal $','loan_on_metal',12),('Melt $','melt',10),('% of melt','pct_melt',9),('Room to 75% $','extra',13),('Stones?','stones',8),
('Last payment','last_payment',12),('Due','due',11),('Pull date','pull',11),('Pawned','pawn',11),('Note',None,28)]
def note(x):
    if not x['melt']: return 'Counter appraisal (coin/bullion/platinum - no weight value)'
    if x['unpriced_items']: return 'Partly counter appraisal (%d item(s) unpriced)'%x['unpriced_items']
    if x['stones']: return 'Stones - check before raising'
    if x['pct_melt'] is not None and x['pct_melt']>=75: return 'Already at/above 75%'
    return ''
wb=Workbook(); S=wb.active; S.title='Summary'
tabs=[('Current','current'),('Grace (past due)','grace (past due, before pull date)'),('Past pull date','past pull date')]
def fill(ws,rows):
    for j,(h,k,w) in enumerate(cols,1):
        c=ws.cell(1,j,h); c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center'); ws.column_dimensions[L(j)].width=w
    rows=sorted(rows,key=lambda x:(-(x['extra'] or 0),x['store']))
    for i,x in enumerate(rows,2):
        for j,(h,k,w) in enumerate(cols,1):
            v=note(x) if k is None else x.get(k)
            if k=='stones': v='YES' if v else ''
            if k=='extra' and not x['melt']: v=None
            if k=='melt' and not x['melt']: v=None
            if k=='pct_melt' and not x['melt']: v=None
            c=ws.cell(i,j,v); c.font=F
            if k in('loan','loan_on_metal','melt','extra'): c.number_format='$#,##0'
            if k=='pct_melt': c.number_format='0"%"'
    ws.freeze_panes='D2'; ws.auto_filter.ref=f"A1:{L(len(cols))}{len(rows)+1}"
for n,s in tabs: fill(wb.create_sheet(n),[x for x in T if x['status']==s])
sh={n:len([x for x in T if x['status']==s]) for n,s in tabs}
S['A1']='Valley Pawn - Precious-metal loan growth list'; S['A1'].font=Font(name='Arial',size=14,bold=True)
S['A2']=f"Paid loans only (at least one payment / extension). Target: raise loan up to 75% of melt. Spot used: gold ${sp['gold']:,.2f}/oz, silver ${sp['silver']:,.2f}/oz ({sp['updated'][:10]}). Built 2026-10-10."; S['A2'].font=F
hd=['Store','Paid PM loans','Priced for melt','Below 75% of melt','Room to 75% $','Avg % of melt (priced)','Counter appraisal','Current','Grace','Past pull date']
for j,h in enumerate(hd,1):
    c=S.cell(4,j,h); c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True); S.column_dimensions[L(j)].width=16
stores=['CUL','HAR','LEX','ROA','WAY']
for i,s in enumerate(stores,5):
    S.cell(i,1,s)
    for tab,col,kind in [('Current','H','c'),('Grace (past due)','I','g'),('Past pull date','J','p')]:
        S[f'{col}{i}']=f"=COUNTIF('{tab}'!A:A,$A{i})"
    S[f'B{i}']=f'=H{i}+I{i}+J{i}'
    t=["'Current'","'Grace (past due)'","'Past pull date'"]
    S[f'C{i}']='='+'+'.join(f"COUNTIFS({x}!$A:$A,$A{i},{x}!$J:$J,\">0\")" for x in t)
    S[f'D{i}']='='+'+'.join(f"COUNTIFS({x}!$A:$A,$A{i},{x}!$L:$L,\">0\")" for x in t)
    S[f'E{i}']='='+'+'.join(f"SUMIFS({x}!$L:$L,{x}!$A:$A,$A{i})" for x in t)
    S[f'F{i}']='=IFERROR(('+'+'.join(f"SUMIFS({x}!$K:$K,{x}!$A:$A,$A{i})" for x in t)+f')/C{i}/100,0)'
    S[f'G{i}']=f'=B{i}-C{i}'
S['A10']='Total'
for col in 'BCDEGHIJ': S[f'{col}10']=f'=SUM({col}5:{col}9)'
S['F10']='=IFERROR(SUMPRODUCT(F5:F9,C5:C9)/C10,0)'
for r in range(5,11):
    for j in range(1,11):
        c=S.cell(r,j); c.font=Font(name='Arial',size=10,bold=(r==10))
    S[f'E{r}'].number_format='$#,##0'; S[f'F{r}'].number_format='0%'
notes=['How to read it: tabs are sorted by Room to 75% (largest first). "Current" = not past due - best call list. "Grace" = past due but before the pull date - mention the payment. "Past pull date" = at risk of forfeit - handle with the store manager.',
'Room to 75% = 75% of melt minus what is already loaned on the metal. Melt = DWT x 0.05 ozt x purity x spot. Coins, bullion and platinum have no weight value in the export, so they show as Counter appraisal.',
'Stones = description mentions stones; the melt ignores stone value, so check at the counter. Rows already at/above 75% are noted, not targets.',
'Data: Bravo Claude Loan Portfolio pulls, ON LOAN items as of 10/10 (5,265 of 5,701 outstanding items, 92%). Three short monthly pulls rerun tonight; list will be refreshed. Nothing has been sent to any customer.']
for i,n in enumerate(notes,12):
    S.cell(i,1,n).font=F; S.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top'); S.merge_cells(start_row=i,start_column=1,end_row=i,end_column=10); S.row_dimensions[i].height=42
wb.save('Valley_Pawn_PM_Loan_Growth_List_20261010.xlsx')
